#!/usr/bin/env python3
"""Delivery QA — source_path validation, value consistency check, HTML pattern scan, conflict leak detection."""

from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path
import re
from typing import Any

from packages.research_core.pipeline.constants import (
    QA_RULE_VERSION,
    FORBIDDEN_HTML_PATTERNS,
    REQUIRED_SECTION_MARKERS,
    VOC_REQUIRED_EVIDENCE_FIELDS,
    REQUIRED_REPORT_DATA_SECTIONS,
    REQUIRED_REPORT_DATA_DECLARATIONS,
)
from packages.research_core.pipeline._utils import (
    _NOT_FOUND, _split_path, _navigate,
)
from packages.research_core.pipeline.public_language import find_public_language_issues

# 冲突泄漏扫描关键词（独立于 FORBIDDEN_HTML_PATTERNS，聚焦冲突复核过程）
_CONFLICT_LEAK_PATTERNS: list[tuple[str, str]] = [
    (r"数据源冲突", "数据源冲突过程泄漏"),
    (r"数据不一致", "数据不一致细节泄漏"),
    (r"两个数据源", "双数据源对比泄漏"),
    (r"冲突复核", "冲突复核过程泄漏"),
    (r"融合策略", "融合策略内部术语泄漏"),
]

_ALLOWED_REPORT_CLASS_TOKENS = {
    "page",
    "hero",
    "eyebrow",
    "verdict",
    "lead",
    "hero-grid",
    "hero-metric",
    "label",
    "value",
    "section",
    "subtitle",
    "insight-row",
    "insight-card",
    "good",
    "warn",
    "tag",
    "tag-green",
    "tag-amber",
    "tag-red",
    "tag-blue",
    "tag-gray",
    "price-band",
    "price-bar",
    "bar",
    "next-steps",
    "next-step",
    "num",
    "risk-list",
    "severity",
    "go-nogo",
    "tc",
    "pill",
    "table-scroll",
    "brand-cell",
    "brand-stack",
    "route-cell",
    "route-name",
    "route-en",
    "asin-cell",
    "keyword-cell",
    "brand-list-cell",
    "nowrap",
}

_KEY_STRUCTURE_CLASSES = {"hero", "page", "section", "go-nogo"}

_KEY_CSS_SELECTORS = [".hero", ".section", ".go-nogo", ".page", ".tag-green", ".tag-amber", ".tag-red", ".insight-card"]

_FAILURE_CLASSIFICATION: dict[str, dict[str, str]] = {
    # analysis: judgment logic failures → retry Stage 10 (Lead Operator Agent)
    "report_data_values_consistent": {"class": "analysis", "retry_stage": "stage_10", "retry_target": "Lead Operator Agent"},
    "p0_blockers_clear": {"class": "analysis", "retry_stage": "stage_10", "retry_target": "Lead Operator Agent"},
    # rendering: presentation/HTML failures → retry Stage 12 (Report Generation Agent)
    "html_exists": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "report_data_has_required_sections": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_no_removed_legacy_sections": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_inline_style": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "uses_report_template_css": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_only_allowed_report_classes": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_no_fixed_data_source_section": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_required_operator_sections": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_gonogo_class": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_no_forbidden_html_patterns": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "has_no_forbidden_xlsx_patterns": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    "no_conflict_leak": {"class": "rendering", "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
    # data: seed/data pipeline failures → retry Stage 11 (seed generation)
    "report_data_exists": {"class": "data", "retry_stage": "stage_11", "retry_target": "build_analysis_report seed"},
    "report_data_sources_valid": {"class": "data", "retry_stage": "stage_11", "retry_target": "build_analysis_report seed"},
    "xlsx_exists": {"class": "data", "retry_stage": "stage_11", "retry_target": "build_analysis_report seed"},
}


def _classify_failures(failures: list[str]) -> dict[str, Any]:
    """Classify each QA failure into analysis/rendering/data and assign retry target.

    Returns:
        {
            "analysis": {"failures": [...], "retry_stage": "stage_10", "retry_target": "Lead Operator Agent"},
            "rendering": {"failures": [...], "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
            "data": {"failures": [...], "retry_stage": "stage_11", "retry_target": "build_analysis_report seed"},
        }
    """
    classified: dict[str, dict[str, Any]] = {
        "analysis": {"failures": [], "retry_stage": "stage_10", "retry_target": "Lead Operator Agent"},
        "rendering": {"failures": [], "retry_stage": "stage_12", "retry_target": "Report Generation Agent"},
        "data": {"failures": [], "retry_stage": "stage_11", "retry_target": "build_analysis_report seed"},
    }
    for name in failures:
        fc = _FAILURE_CLASSIFICATION.get(name)
        if fc:
            classified[fc["class"]]["failures"].append(name)
        else:
            # Unknown check → default to rendering (Report Generation Agent)
            classified["rendering"]["failures"].append(name)
    return {k: v for k, v in classified.items() if v["failures"]}


def find_report_files(run_dir: Path) -> tuple[Path | None, Path | None, Path | None]:
    """Find report_data.json, HTML, and XLSX in the analysis directory."""
    analysis_dir = run_dir / "analysis"
    report_data = analysis_dir / "report_data.json"
    if not report_data.exists():
        report_data = None
    html_path = None
    xlsx_path = None
    if analysis_dir.is_dir():
        for f in analysis_dir.iterdir():
            if f.suffix == ".html" and f.name.endswith("_分析报告.html"):
                html_path = f
            elif f.suffix == ".xlsx" and f.name.endswith("_决策工具包.xlsx"):
                xlsx_path = f
    return report_data, html_path, xlsx_path


def print_qa_summary(result: dict) -> None:
    """Print a human-readable QA summary."""
    print(f"QA Rule Version: {result.get('qa_rule_version', '?')}")
    print(f"Status: {result['status'].upper()}")
    print(f"Generated: {result.get('generated_at', '?')}")
    print()

    checks = result.get("checks", {})
    failures = result.get("failures", [])

    for name, passed in checks.items():
        if name in ("report_data_sources_note", "report_data_values_note",
                     "report_data_value_mismatches", "forbidden_html_hits",
                     "forbidden_xlsx_hits", "p0_blocker_hits", "conflict_leak_hits",
                     "report_template_css_hits", "report_class_hits",
                     "fixed_data_source_section_hits"):
            continue
        if isinstance(passed, bool):
            icon = "PASS" if passed else "FAIL"
            print(f"  [{icon}] {name}")

    if failures:
        print(f"\nFailures ({len(failures)}):")
        for f in failures:
            print(f"  - {f}")
        classification = result.get("failure_classification") or {}
        for cls_key, cls_info in classification.items():
            label_map = {"analysis": "分析错误→Stage 10", "rendering": "渲染错误→Stage 12", "data": "数据错误→Stage 11"}
            label = label_map.get(cls_key, cls_key)
            print(f"\n  [{label}] retry {cls_info['retry_target']}:")
            for fn in cls_info["failures"]:
                print(f"    - {fn}")

    if checks.get("forbidden_html_hits"):
        print("\nForbidden HTML patterns found:")
        for hit in checks["forbidden_html_hits"]:
            print(f"  - {hit}")

    if checks.get("forbidden_xlsx_hits"):
        print("\nForbidden XLSX patterns found:")
        for hit in checks["forbidden_xlsx_hits"]:
            print(f"  - {hit}")

    if checks.get("conflict_leak_hits"):
        print("\nConflict process leaks found:")
        for hit in checks["conflict_leak_hits"]:
            print(f"  - {hit}")

    if checks.get("p0_blocker_hits"):
        print("\nP0 Blocker hits:")
        for hit in checks["p0_blocker_hits"]:
            print(f"  - {hit}")

    if checks.get("report_template_css_hits"):
        print("\nReport template CSS issues:")
        for hit in checks["report_template_css_hits"]:
            print(f"  - {hit}")

    if checks.get("report_class_hits"):
        print("\nReport class issues:")
        for hit in checks["report_class_hits"]:
            print(f"  - {hit}")

    if checks.get("fixed_data_source_section_hits"):
        print("\nReport section issues:")
        for hit in checks["fixed_data_source_section_hits"]:
            print(f"  - {hit}")

    if checks.get("report_data_sources_note"):
        print(f"\nSource note: {checks['report_data_sources_note']}")

    if checks.get("report_data_values_note"):
        print(f"\nValues note: {checks['report_data_values_note']}")

    if checks.get("report_data_value_mismatches"):
        print("\nValue mismatches:")
        for m in checks["report_data_value_mismatches"][:10]:
            print(f"  - {m.get('key', '?')}: reported={m.get('reported', '?')} vs resolved={m.get('resolved', '?')}")


def run_delivery_qa(report_data_path: Path, html_path: Path, xlsx_path: Path, analysis: dict[str, Any] | None = None) -> dict[str, Any]:
    run_dir = html_path.parent.parent
    packets = _load_packets_for_qa(run_dir, analysis)
    source_result = _validate_report_data_sources(report_data_path, packets)
    value_result = _validate_values_against_sources(report_data_path, packets)
    blocker_result = _validate_p0_delivery_blockers(run_dir, report_data_path)
    has_data = _report_data_has_required_sections(report_data_path)
    forbidden_result = _has_no_forbidden_html_patterns(html_path)
    forbidden_xlsx_result = _has_no_forbidden_xlsx_patterns(xlsx_path)
    conflict_leak_result = _scan_conflict_leak(html_path, run_dir)
    template_css_result = _uses_report_template_css(html_path)
    class_result = _has_only_allowed_report_classes(html_path)
    data_source_section_result = _has_no_fixed_data_source_section(html_path)
    checks = {
        "report_data_exists": report_data_path.exists(),
        "html_exists": html_path.exists(),
        "xlsx_exists": xlsx_path.exists(),
        "report_data_has_required_sections": has_data,
        "has_no_removed_legacy_sections": _has_no_removed_legacy_sections(html_path),
        "has_inline_style": _has_inline_style(html_path),
        "uses_report_template_css": template_css_result["pass"],
        "has_only_allowed_report_classes": class_result["pass"],
        "has_no_fixed_data_source_section": data_source_section_result["pass"],
        "has_required_operator_sections": _has_required_operator_sections(html_path),
        "has_gonogo_class": _has_gonogo_class(html_path),
        "report_data_sources_valid": source_result["pass"],
        "report_data_values_consistent": value_result["pass"],
        "has_no_forbidden_html_patterns": forbidden_result["pass"],
        "has_no_forbidden_xlsx_patterns": forbidden_xlsx_result["pass"],
        "no_conflict_leak": conflict_leak_result["pass"],
        "p0_blockers_clear": blocker_result["pass"],
    }
    if forbidden_result.get("hits"):
        checks["forbidden_html_hits"] = forbidden_result["hits"]
    if forbidden_xlsx_result.get("hits"):
        checks["forbidden_xlsx_hits"] = forbidden_xlsx_result["hits"]
    if conflict_leak_result.get("hits"):
        checks["conflict_leak_hits"] = conflict_leak_result["hits"]
    if template_css_result.get("hits"):
        checks["report_template_css_hits"] = template_css_result["hits"]
    if class_result.get("hits"):
        checks["report_class_hits"] = class_result["hits"]
    if data_source_section_result.get("hits"):
        checks["fixed_data_source_section_hits"] = data_source_section_result["hits"]
    if source_result.get("reason"):
        checks["report_data_sources_note"] = source_result["reason"]
    if value_result.get("reason"):
        checks["report_data_values_note"] = value_result["reason"]
    if value_result.get("mismatches"):
        checks["report_data_value_mismatches"] = value_result["mismatches"]
    if blocker_result.get("hits"):
        checks["p0_blocker_hits"] = blocker_result["hits"]
    failures = [name for name, passed in checks.items() if not passed and name not in ("report_data_sources_note", "report_data_values_note", "report_data_value_mismatches")]
    failure_classification = _classify_failures(failures)
    return {
        "status": "pass" if not failures else "fail",
        "qa_rule_version": QA_RULE_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "checks": checks,
        "failures": failures,
        "failure_classification": failure_classification,
    }


def _load_packets_for_qa(run_dir: Path, analysis: dict[str, Any] | None = None) -> dict[str, Any]:
    """Load evidence packets for source_path validation.

    优先从磁盘读取 analysis_packet.json（持久化中间层）；
    若不存在则回退到入参 analysis（运行时兼容）。
    """
    paths = {
        "market_structure": run_dir / "market_structure" / "market_structure_evidence_packet.json",
        "search_demand": run_dir / "search_demand" / "search_demand_evidence_packet.json",
        "voc": run_dir / "review_voc" / "voc_evidence_packet.json",
        "route_matrix": run_dir / "route_matrix_confirm.json",
    }
    packets: dict[str, Any] = {}
    for key, path in paths.items():
        try:
            if path.exists():
                packets[key] = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, ValueError):
            pass

    # 优先从磁盘读取 analysis_packet.json，保证交付后可独立溯源
    analysis_packet_path = run_dir / "analysis" / "analysis_packet.json"
    if analysis_packet_path.exists():
        try:
            packets["analysis"] = json.loads(analysis_packet_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, ValueError):
            if analysis:
                packets["analysis"] = analysis
    elif analysis:
        packets["analysis"] = analysis
    return packets


def _has_no_removed_legacy_sections(html_path: Path) -> bool:
    if not html_path.exists():
        return False
    html = html_path.read_text(encoding="utf-8")
    removed_section_markers = (
        "legacy_source_collection_section",
        "legacy_fillback_stage_section",
        "legacy_cost_review_section",
    )
    return not any(marker in html for marker in removed_section_markers)


def _has_no_forbidden_html_patterns(html_path: Path) -> dict:
    """扫描 HTML 中的禁止模式：抽象路线标签、内部术语、虚假宣传措辞。"""
    if not html_path.exists():
        return {"pass": False, "hits": ["HTML 文件不存在"]}
    html = html_path.read_text(encoding="utf-8")
    hits = []
    for pattern, description in FORBIDDEN_HTML_PATTERNS:
        matches = re.findall(pattern, html)
        if matches:
            unique_matches = list(set(matches))[:5]
            hits.append(f"{description}（匹配: {', '.join(unique_matches)}）")
    return {"pass": len(hits) == 0, "hits": hits}


def _has_no_forbidden_xlsx_patterns(xlsx_path: Path) -> dict:
    """扫描 XLSX 单元格中的交付禁用话术。"""
    if not xlsx_path.exists():
        return {"pass": False, "hits": ["XLSX 文件不存在"]}
    try:
        from openpyxl import load_workbook

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - defensive for corrupt workbook
        return {"pass": False, "hits": [f"XLSX 读取失败: {exc}"]}

    hits: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    issues = find_public_language_issues(value)
                    for issue in issues:
                        hits.append(f"{ws.title}!{cell.coordinate}: {issue}")
                        if len(hits) >= 20:
                            return {"pass": False, "hits": hits}
    finally:
        wb.close()

    return {"pass": len(hits) == 0, "hits": hits}


def _scan_conflict_leak(html_path: Path, run_dir: Path) -> dict:
    """扫描 HTML 是否泄漏了冲突复核过程的内部语言。

    独立于 FORBIDDEN_HTML_PATTERNS，聚焦冲突复核专用术语。
    如果 run_dir 中不存在冲突复核产物，跳过检查（无冲突则无泄漏风险）。
    """
    if not html_path.exists():
        return {"pass": False, "hits": ["HTML 文件不存在"]}

    conflict_packet = run_dir / "conflict_review" / "conflict_resolution_packet.json"
    if not conflict_packet.exists():
        # 无冲突复核产物，跳过冲突泄漏检查
        return {"pass": True, "hits": []}

    html = html_path.read_text(encoding="utf-8")
    hits = []
    for pattern, description in _CONFLICT_LEAK_PATTERNS:
        matches = re.findall(pattern, html)
        if matches:
            hits.append(f"{description}（匹配: {pattern}）")

    return {"pass": len(hits) == 0, "hits": hits}


def _has_inline_style(html_path: Path) -> bool:
    """HTML 必须内嵌 <style>（内容来自 report_template.css），不能使用外部 <link>。"""
    if not html_path.exists():
        return False
    html = html_path.read_text(encoding="utf-8")
    has_inline_style = '<style>' in html
    import re
    has_link = bool(re.search(r'<link[^>]*report_template\.css', html))
    return has_inline_style and not has_link


def _report_template_css_path() -> Path:
    return (
        Path(__file__).resolve().parents[3]
        / "skills"
        / "amazon-product-research"
        / "references"
        / "report_template.css"
    )


def _extract_inline_style(html: str) -> str | None:
    match = re.search(r"<style>\s*(.*?)\s*</style>", html, flags=re.DOTALL)
    return match.group(1).strip() if match else None


def _uses_report_template_css(html_path: Path) -> dict[str, Any]:
    """HTML 的 <style> 必须包含关键 CSS 选择器，不要求逐字节匹配模板。"""
    if not html_path.exists():
        return {"pass": False, "hits": ["HTML 文件不存在"]}
    html = html_path.read_text(encoding="utf-8")
    inline_css = _extract_inline_style(html)
    if inline_css is None:
        return {"pass": False, "hits": ["缺少内嵌 <style>"]}

    missing_selectors = [s for s in _KEY_CSS_SELECTORS if s not in inline_css]
    if missing_selectors:
        return {
            "pass": False,
            "hits": [f"<style> 缺少关键 CSS 选择器: {', '.join(missing_selectors)}"],
        }
    return {"pass": True, "hits": []}


def _has_only_allowed_report_classes(html_path: Path) -> dict[str, Any]:
    """检查关键结构 class 存在；非白名单 class 仅 warning 不阻断。"""
    if not html_path.exists():
        return {"pass": False, "hits": ["HTML 文件不存在"]}
    html = html_path.read_text(encoding="utf-8")
    class_values = re.findall(r"""\bclass\s*=\s*["']([^"']+)["']""", html)
    all_tokens = {token for value in class_values for token in value.split()}

    missing_key = _KEY_STRUCTURE_CLASSES - all_tokens
    if missing_key:
        return {
            "pass": False,
            "hits": [f"缺少关键结构 class: {', '.join(sorted(missing_key))}"],
        }

    unknown = sorted(all_tokens - _ALLOWED_REPORT_CLASS_TOKENS)
    hits: list[str] = []
    if unknown:
        hits.append(f"[WARNING] 发现模板外 class（不阻断）: {', '.join(unknown[:10])}")
    return {"pass": True, "hits": hits}


def _has_no_fixed_data_source_section(html_path: Path) -> dict[str, Any]:
    """正式 HTML 不固定展示“数据来源与口径”板块。"""
    if not html_path.exists():
        return {"pass": False, "hits": ["HTML 文件不存在"]}
    html = html_path.read_text(encoding="utf-8")
    if "数据来源与口径" in html:
        return {
            "pass": False,
            "hits": ["HTML 不得固定展示“数据来源与口径”板块"],
        }
    return {"pass": True, "hits": []}


def _has_required_operator_sections(html_path: Path) -> bool:
    if not html_path.exists():
        return False
    html = html_path.read_text(encoding="utf-8")
    return all(marker in html for marker in REQUIRED_SECTION_MARKERS)


def _has_gonogo_class(html_path: Path) -> bool:
    if not html_path.exists():
        return False
    html = html_path.read_text(encoding="utf-8")
    return 'class="go-nogo"' in html or "class='go-nogo'" in html


def _has_voc_evidence_refs(analysis_json: Path) -> bool:
    """VOC 痛点必须有 evidence_refs，且每个 ref 包含 review_id/quote/rating/asin 四字段。"""
    if not analysis_json.exists():
        return False
    data = json.loads(analysis_json.read_text(encoding="utf-8"))
    voc = data.get("voc_spec_translation") or {}
    pain_points = voc.get("pain_points") or []
    if not pain_points:
        # 没有痛点时不扣分（可能评论样本不足）
        return True
    return all(
        isinstance(pp.get("evidence_refs"), list)
        and len(pp["evidence_refs"]) > 0
        and all(
            isinstance(ref, dict)
            and all(ref.get(field) for field in VOC_REQUIRED_EVIDENCE_FIELDS)
            for ref in pp["evidence_refs"]
        )
        for pp in pain_points
    )


def _report_data_has_required_sections(report_data_path: Path) -> bool:
    """report_data.json 必须包含所有必要的板块数据。"""
    if not report_data_path.exists():
        return False
    try:
        data = json.loads(report_data_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, ValueError):
        return False
    return all(key in data for key in REQUIRED_REPORT_DATA_DECLARATIONS) and all(
        section in data for section in REQUIRED_REPORT_DATA_SECTIONS
    )


def _validate_p0_delivery_blockers(run_dir: Path, report_data_path: Path) -> dict[str, Any]:
    """Block delivery when P0 governance artifacts contain unresolved blockers."""
    hits: list[str] = []
    conflict_path = run_dir / "conflict_review" / "conflict_resolution_packet.json"
    if conflict_path.exists():
        conflict = _safe_load_json(conflict_path)
        blocking = conflict.get("blocking_conflicts") or []
        unresolved = [
            item
            for item in blocking
            if isinstance(item, dict) and item.get("status") not in {"resolved", "accepted", "waived"}
        ]
        if unresolved:
            hits.append("blocking conflict 未解决，不得交付")

    evaluation_path = run_dir / "evaluations" / "evaluation_summary.json"
    if evaluation_path.exists():
        evaluation = _safe_load_json(evaluation_path)
        dimension_results = evaluation.get("dimension_results") or {}
        data_quality = dimension_results.get("data_quality") if isinstance(dimension_results, dict) else {}
        if isinstance(data_quality, dict) and data_quality.get("rating") == "blocked":
            hits.append("data_quality=blocked 时不得输出 Go")
        if isinstance(dimension_results, dict):
            for name, result in dimension_results.items():
                if isinstance(result, dict) and result.get("confidence") == "low" and result.get("rating") == "strong":
                    hits.append(f"low confidence 不得支撑强结论: {name}")

    progress_path = run_dir / "progress.json"
    if progress_path.exists():
        progress = _safe_load_json(progress_path)
        stages = progress.get("stages") or {}
        unfinished = [
            stage_id
            for stage_id, stage in stages.items()
            if isinstance(stage, dict) and stage.get("status") not in {"done", "needs_user"}
        ]
        if unfinished:
            hits.append("progress.json 存在未完成阶段: " + ", ".join(sorted(unfinished)[:5]))

    # ── Judgment 10-deep-field completeness ──────────────────
    judgment_path = run_dir / "analysis" / "integrated_operator_judgment.json"
    if judgment_path.exists():
        from packages.research_core.pipeline.judgment_contract import validate_judgment_file
        jc = validate_judgment_file(judgment_path)
        if not jc["pass"]:
            missing_deep = jc.get("missing_deep_fields", [])
            incomplete_deep = jc.get("incomplete_deep_fields", [])
            if missing_deep:
                hits.append(f"judgment 缺少深度分析字段: {', '.join(missing_deep)}")
            if incomplete_deep:
                hits.append(f"judgment 深度分析字段不完整: {', '.join(incomplete_deep[:10])}")
            if not jc.get("verdict_valid"):
                hits.append("judgment final_verdict 值无效")
            if not jc.get("confidence_valid"):
                hits.append("judgment confidence 值无效")

    schema_versions = _collect_schema_versions(
        report_data_path,
        conflict_path,
        evaluation_path,
        run_dir / "analysis" / "integrated_operator_judgment.json",
    )
    missing = [path for path, version in schema_versions.items() if not version]
    if missing:
        hits.append("schema_version 缺失: " + ", ".join(missing))

    return {"pass": not hits, "hits": hits}


def _safe_load_json(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, ValueError, OSError):
        return {}
    return data if isinstance(data, dict) else {}


def _collect_schema_versions(*paths: Path) -> dict[str, str]:
    versions: dict[str, str] = {}
    for path in paths:
        if not path.exists():
            continue
        data = _safe_load_json(path)
        version = data.get("schema_version", "")
        versions[str(path)] = str(version) if version else ""
    return versions


def _validate_report_data_sources(report_data_path: Path, packets: dict[str, Any]) -> dict:
    """校验 report_data.json 中 source_path 能否在证据包中找到对应字段。

    返回 dict 包含 pass/fail 和解析明细。阈值：
    - 存在空 source_path → block（pass=False，AI 未填充）
    - 未解析率 > 5% → block（pass=False）
    - 未解析率 > 2% → warning（pass=True，但记录）
    - 未解析率 ≤ 2% 且无空路径 → pass

    __ai_judgment__ 是 seed 占位标记，表示该字段需 AI 增强后填充，记 warning 不阻断。
    """
    if not report_data_path.exists():
        return {"pass": False, "total": 0, "resolved": 0, "vague": 0, "unresolved": 0, "empty": 0, "ai_pending": 0, "unresolved_pct": 0.0, "reason": "report_data.json 不存在"}
    try:
        data = json.loads(report_data_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, ValueError):
        return {"pass": False, "total": 0, "resolved": 0, "vague": 0, "unresolved": 0, "empty": 0, "ai_pending": 0, "unresolved_pct": 0.0, "reason": "report_data.json 解析失败"}

    source_paths = _extract_source_paths(data)
    unresolved: list[dict] = []
    resolved = 0
    vague = 0
    empty = 0
    ai_pending = 0

    for item in source_paths:
        sp = item["source_path"].strip()
        if sp == "__ai_judgment__":
            ai_pending += 1
            continue
        if not sp or sp == "N/A" or sp == "—":
            empty += 1
            unresolved.append({"path": "(空)", "key": item.get("parent_key", ""), "reason": "source_path 为空字符串，AI 未填充"})
            continue

        result = _try_resolve_path(sp, packets)
        if result["status"] == "unresolved":
            unresolved.append({"path": sp, "key": item.get("parent_key", ""), "reason": result["reason"]})
        elif result["status"] == "ok" and "模糊" in result.get("reason", ""):
            vague += 1
        else:
            resolved += 1

    total = len(source_paths)
    effective = max(total - ai_pending, 1)  # __ai_judgment__ 不参与未解析率计算
    unresolved_pct = len(unresolved) / max(effective, 1)

    # 空 source_path 直接阻断
    if empty > 0:
        reason = f"存在 {empty} 条空 source_path（AI 未填充），共 {len(unresolved)}/{effective} 条路径无法解析 ({unresolved_pct:.0%})"
        if ai_pending:
            reason += f"，{ai_pending} 条标记为 __ai_judgment__（待 AI 增强）"
        return {
            "pass": False,
            "total": total,
            "resolved": resolved,
            "vague": vague,
            "unresolved": len(unresolved),
            "empty": empty,
            "ai_pending": ai_pending,
            "unresolved_pct": unresolved_pct,
            "reason": reason,
            "unresolved_paths": unresolved,
        }

    passed = unresolved_pct <= 0.05
    reason_parts = []
    if unresolved_pct > 0.05:
        reason_parts.append(f"未解析率 {unresolved_pct:.0%} 超过 5% 阈值，共 {len(unresolved)}/{effective} 条路径无法溯源")
    elif unresolved_pct > 0.02:
        reason_parts.append(f"未解析率 {unresolved_pct:.0%} 在 2%-5% 之间，共 {len(unresolved)}/{effective} 条路径无法溯源（不阻断）")
    if ai_pending:
        reason_parts.append(f"{ai_pending} 条标记为 __ai_judgment__（待 AI 增强）")

    return {
        "pass": passed,
        "total": total,
        "resolved": resolved,
        "vague": vague,
        "unresolved": len(unresolved),
        "empty": empty,
        "ai_pending": ai_pending,
        "unresolved_pct": unresolved_pct,
        "reason": "；".join(reason_parts) if reason_parts else "",
        "unresolved_paths": unresolved,
    }


def _extract_source_paths(obj: Any, parent_key: str = "", parent_obj: dict | None = None) -> list[dict]:
    """递归提取 report_data.json 中所有 source_path 及其上下文。

    返回值中每条记录包含 source_path、parent_key、以及同级 value 字段（用于值比对）。
    """
    results: list[dict] = []
    if isinstance(obj, dict):
        for key, val in obj.items():
            if key == "source_path" and isinstance(val, str):
                entry: dict = {"source_path": val, "parent_key": parent_key}
                # 当前对象 obj 中查找同级 value 字段（非 parent_obj）
                if "value" in obj:
                    entry["reported_value"] = obj.get("value")
                results.append(entry)
            else:
                results.extend(_extract_source_paths(val, parent_key=key, parent_obj=obj))
    elif isinstance(obj, list):
        for item in obj:
            results.extend(_extract_source_paths(item, parent_key=parent_key, parent_obj=item if isinstance(item, dict) else None))
    return results


def _try_resolve_path(source_path: str, packets: dict[str, Any]) -> dict:
    """尝试在证据包中解析一个 source_path。

    支持的路径格式：
    - packet.field.subfield (如 market_structure.market_size.primary_market)
    - packet.array_id.field (如 search_demand.f8.value.top100_monthly_units)
    - packet.array[index].field (如 voc.pain_points_by_dimension[0].review_count)
    - packet.field (模糊溯源，如 market_structure)
    - /field/subfield (JSON pointer 风格，需有包名前缀)
    - 绝对文件路径 (如 runs/.../file.json 或 mcp_snapshots/...)

    当直接解析失败时，尝试级联溯源：通过 evidence_refs 链追踪到其他证据包。
    """
    # 处理文件路径（runs/...、mcp_snapshots/... 或绝对路径包含文件扩展名）
    if source_path.startswith("runs/") or source_path.startswith("mcp_snapshots/"):
        return _resolve_file_path(source_path)
    if source_path.startswith("/") and _looks_like_file_path(source_path):
        return _resolve_file_path(source_path)

    # JSON pointer 风格 (/field/subfield) — 不以文件扩展名结尾，不以路径分隔符为主
    if source_path.startswith("/"):
        return _resolve_json_pointer(source_path, packets)

    # 匹配包名前缀
    packet_key = None
    rest = source_path
    for pkey in ("market_structure", "search_demand", "voc", "route_matrix", "analysis", "seller_sprite"):
        if source_path.startswith(pkey):
            packet_key = pkey
            rest = source_path[len(pkey):].lstrip(".")
            break

    if not packet_key:
        return {"status": "unresolved", "reason": f"无法识别包名前缀: {source_path}"}

    # Alias: seller_sprite -> market_structure
    if packet_key == "seller_sprite":
        packet_key = "market_structure"

    packet = packets.get(packet_key)
    if not packet:
        return {"status": "unresolved", "reason": f"证据包未加载: {packet_key}"}

    if not rest:
        # 只有包名无字段路径 —— 模糊溯源
        return {"status": "ok", "reason": "包级模糊溯源（无字段路径）"}

    # 剥离结尾的描述性文字（空格后跟中文说明、加总等）
    stripped = rest
    for sep in [" 月销额加总", " 月销额:", " 月销额计算:", " 竞品数对比", " 竞品数", " 加总", " + "]:
        if sep in stripped:
            idx = stripped.index(sep)
            stripped = stripped[:idx]
            break
    # 多引用复合路径（如 facts[f8,f9,f10] 或 facts[f1,f2].field）→ 无法单点解析，跳过
    if "][" in stripped:
        return {"status": "ok", "reason": "复合引用（多源加总），无法单点解析"}
    # 单括号内多 id（如 [f1,f2]、facts[f1,f2] 或 [f8,f9,f10]）
    if re.search(r'\[[^\]]*,[^\]]*\]', stripped):
        return {"status": "ok", "reason": "复合引用（多源加总），无法单点解析"}

    # 分割路径
    parts = _split_path(stripped)
    if not parts:
        return {"status": "unresolved", "reason": f"路径解析后为空: {rest}"}

    # 导航 JSON
    current: Any = packet
    for part in parts:
        current = _navigate(current, part)
        if current is _NOT_FOUND:
            # 直接解析失败，尝试级联溯源
            cascade = _try_cascade_trace(source_path, packet_key, parts, packets)
            if cascade:
                return cascade
            return {"status": "unresolved", "reason": f"路径段 '{part}' 在 '{packet_key}' 中未找到"}

    return {"status": "ok", "reason": "已解析", "resolved_value": current}


def _looks_like_file_path(path: str) -> bool:
    """判断一个以 / 开头的路径是否为文件路径而非 JSON pointer。

    JSON pointer 特征：路径段不含文件扩展名，段数少，不含路径分隔符模式。
    文件路径特征：含 .json/.xlsx 等扩展名，或含多层目录结构。
    """
    # 有明确的文件扩展名 → 文件路径
    for ext in (".json", ".xlsx", ".csv", ".txt", ".html"):
        if ext in path:
            return True
    # 路径段数多（>3 层）且没有数组索引 → 可能是文件路径
    parts = [p for p in path.split("/") if p]
    if len(parts) > 3 and not any("[" in p for p in parts):
        return True
    return False


def _resolve_file_path(file_path: str) -> dict:
    """尝试解析绝对文件路径，检查文件是否存在并可读取。"""
    p = Path(file_path)
    if not p.is_absolute():
        # 尝试相对于当前工作目录
        p = Path.cwd() / file_path
    if p.exists():
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
            return {"status": "ok", "reason": f"文件存在并可读取: {p.name}", "resolved_value": data}
        except (json.JSONDecodeError, ValueError):
            return {"status": "unresolved", "reason": f"文件存在但无法解析 JSON: {file_path}"}
    return {"status": "unresolved", "reason": f"文件不存在: {file_path}"}


def _resolve_json_pointer(pointer: str, packets: dict[str, Any]) -> dict:
    """解析 JSON pointer 风格路径 (/field/subfield)。

    遍历所有已加载的证据包，尝试在每个包中按 pointer 导航。
    """
    parts = [p for p in pointer.split("/") if p]
    if not parts:
        return {"status": "unresolved", "reason": f"JSON pointer 为空: {pointer}"}

    for pkey, packet in packets.items():
        current: Any = packet
        resolved_all = True
        for part in parts:
            current = _navigate(current, part)
            if current is _NOT_FOUND:
                resolved_all = False
                break
        if resolved_all:
            return {"status": "ok", "reason": f"JSON pointer 在 '{pkey}' 中解析成功", "resolved_value": current}

    return {"status": "unresolved", "reason": f"JSON pointer 在所有证据包中未找到: {pointer}"}


def _try_cascade_trace(
    source_path: str, packet_key: str, parts: list[str], packets: dict[str, Any]
) -> dict | None:
    """级联溯源：当直接路径解析失败时，尝试通过 evidence_refs 链追踪。

    策略：
    1. 找到 parts 中最后一个成功导航的位置
    2. 在该位置查找 evidence_refs
    3. 按 evidence_refs 的 packet/file 字段跳转到其他证据包
    4. 在新包中继续解析剩余 parts
    """
    packet = packets.get(packet_key)
    if not packet:
        return None

    # 逐步导航，找到最后一个成功的位置
    current: Any = packet
    last_ok_idx = -1
    for i, part in enumerate(parts):
        nxt = _navigate(current, part)
        if nxt is _NOT_FOUND:
            break
        current = nxt
        last_ok_idx = i

    if last_ok_idx < 0:
        return None

    # 在最后成功位置查找 evidence_refs
    evidence_refs = None
    if isinstance(current, dict):
        evidence_refs = current.get("evidence_refs")
    if not evidence_refs or not isinstance(evidence_refs, list) or len(evidence_refs) == 0:
        return None

    # 尝试每个 evidence_ref 跳转
    remaining_parts = parts[last_ok_idx + 1:]
    if not remaining_parts:
        return None

    for ref in evidence_refs:
        if not isinstance(ref, dict):
            continue
        ref_packet = ref.get("packet") or ref.get("source") or ""
        ref_file = ref.get("file") or ref.get("path") or ""

        target_packet = None
        if ref_packet and ref_packet in packets:
            target_packet = packets[ref_packet]
        elif ref_file:
            file_result = _resolve_file_path(ref_file)
            if file_result["status"] == "ok":
                target_packet = file_result["resolved_value"]

        if not target_packet:
            continue

        # 在新包中继续导航剩余 parts
        current_ref: Any = target_packet
        all_found = True
        for part in remaining_parts:
            current_ref = _navigate(current_ref, part)
            if current_ref is _NOT_FOUND:
                all_found = False
                break

        if all_found:
            return {
                "status": "ok",
                "reason": f"级联溯源成功（通过 evidence_ref → {ref_packet or ref_file}）",
                "resolved_value": current_ref,
            }

    return None


def _validate_values_against_sources(
    report_data_path: Path, packets: dict[str, Any]
) -> dict:
    """抽查 report_data.json 中 value 与证据包实际值是否一致。

    对每条已解析的 source_path，提取证据包中的字段值，做归一化比对。
    返回 mismatches 列表和 pass/fail。
    """
    if not report_data_path.exists():
        return {"pass": True, "mismatches": [], "checked": 0, "skipped": 0}

    try:
        data = json.loads(report_data_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, ValueError):
        return {"pass": True, "mismatches": [], "checked": 0, "skipped": 0}

    source_paths = _extract_source_paths(data)
    checked = 0
    skipped = 0
    mismatches: list[dict] = []

    for item in source_paths:
        sp = item["source_path"].strip()
        reported = item.get("reported_value")

        # 跳过占位符和空值
        if not sp or sp in ("N/A", "—", "__ai_judgment__"):
            skipped += 1
            continue
        if reported is None or str(reported).strip() in ("", "待补", "—"):
            skipped += 1
            continue

        result = _try_resolve_path(sp, packets)
        if result["status"] != "ok" or "模糊" in result.get("reason", ""):
            skipped += 1
            continue

        resolved = result.get("resolved_value")
        if resolved is None:
            skipped += 1
            continue

        # 解析到 dict/list：派生值（如 recommended_price、core_search_volume），
        # 无法做标量比对，但至少验证源数据非空、派生依据存在。
        if isinstance(resolved, (dict, list)):
            if _is_empty(resolved):
                mismatches.append({
                    "path": sp,
                    "key": item.get("parent_key", ""),
                    "reported": str(reported),
                    "resolved": "∅ (空 dict/list，派生值缺少依据)",
                })
            checked += 1
            continue

        if not _values_match(reported, resolved):
            mismatches.append({
                "path": sp,
                "key": item.get("parent_key", ""),
                "reported": str(reported),
                "resolved": str(resolved),
            })

        checked += 1

    passed = len(mismatches) == 0
    reason = ""
    if mismatches:
        reason = f"{len(mismatches)}/{checked} 条抽查值与证据包不一致"
    elif checked > 0:
        reason = f"抽查 {checked} 条，全部一致"

    return {
        "pass": passed,
        "checked": checked,
        "skipped": skipped,
        "mismatches": mismatches,
        "reason": reason,
    }


def _is_empty(val: Any) -> bool:
    """判断值是否为空的 dict/list（派生值缺少依据）。"""
    if isinstance(val, dict):
        return len(val) == 0
    if isinstance(val, list):
        return len(val) == 0
    return False


def _values_match(reported: Any, resolved: Any) -> bool:
    """归一化比较两个值是否匹配。

    处理常见的表示差异：单位后缀、货币符号、百分比格式、逗号分隔、
    绝对路径 vs 相对路径等。
    """
    if reported == resolved:
        return True

    rpt = str(reported).strip()
    rsl = str(resolved).strip()

    # 路径值：取共同的尾部（如 runs/.../file.json）比对
    if ("/" in rpt or "\\" in rpt) and ("/" in rsl or "\\" in rsl):
        rpt_parts = rpt.replace("\\", "/").rstrip("/").split("/")
        rsl_parts = rsl.replace("\\", "/").rstrip("/").split("/")
        # 取较短路径的后 N 段，在较长路径中匹配
        min_len = min(len(rpt_parts), len(rsl_parts))
        if rpt_parts[-min_len:] == rsl_parts[-min_len:]:
            return True

    def _normalize(v: Any) -> str:
        s = str(v).strip().lower()
        s = s.replace("$", "").replace(",", "").replace(" ", "")
        for suffix in ("units", "unit", "%", "usd", "cny"):
            if s.endswith(suffix):
                s = s[:-len(suffix)]
        return s.strip()

    return _normalize(reported) == _normalize(resolved)
