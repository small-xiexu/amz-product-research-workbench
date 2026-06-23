"""Excel reading utilities for SellerSprite manual export files."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required to read Excel exports") from exc


# ---------------------------------------------------------------------------
# Type coercion
# ---------------------------------------------------------------------------

def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def to_float(value: Any) -> float | None:
    if value in (None, "", "--", "--,--"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("$", "").replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def to_int(value: Any) -> int | None:
    number = to_float(value)
    return int(number) if number is not None else None


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def slugify(text: str) -> str:
    lowered = text.lower()
    slug = re.sub(r"[^a-z0-9]+", "-", lowered).strip("-")
    return slug or "manual-export-market"


def contains_chinese(text: str) -> bool:
    return bool(re.search(r"[一-鿿]", text))


def display_keyword(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text.replace("-", " ")).strip()
    if not normalized or contains_chinese(normalized):
        return normalized
    small_words = {"and", "or", "with", "for", "of", "in", "to", "the", "a", "an"}
    words = []
    for index, word in enumerate(normalized.split(" ")):
        lower = word.lower()
        words.append(lower if index > 0 and lower in small_words else lower.capitalize())
    return " ".join(words)


def clean_task_name(text: Any) -> str:
    value = str(text or "").strip()
    if not value:
        return ""
    internal_markers = [
        "UI重构预览", "UI 重构预览", "UI重构版", "UI 重构版",
        "重构预览", "预览版", "重新验证", "重新驗證",
    ]
    for marker in internal_markers:
        value = value.replace(marker, "")
    value = re.sub(r"[_｜|/\\-]+$", "", value)
    value = re.sub(r"[_｜|/\\-]{2,}", "_", value)
    return value.strip(" _-/｜|\\")


def derive_market_name(manifest: dict[str, Any], seed_keyword: str) -> str:
    task_name = clean_task_name(manifest.get("metadata", {}).get("task_name"))
    seed_name = display_keyword(seed_keyword)
    if task_name and seed_name:
        if contains_chinese(task_name) and not contains_chinese(seed_name):
            return f"{task_name} / {seed_name}"
        if task_name.lower() != seed_name.lower():
            return task_name
    return task_name or seed_name or "Manual Export Market"


# ---------------------------------------------------------------------------
# Excel reading
# ---------------------------------------------------------------------------

def read_records(xlsx_path: Path, sheet_name: str, header_row: int) -> list[dict[str, Any]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    header_values = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    headers = [str(clean_cell(v)) if clean_cell(v) not in (None, "") else "" for v in header_values]
    records: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = [clean_cell(v) for v in row]
        if not any(v not in (None, "") for v in values):
            continue
        records.append({h: values[i] for i, h in enumerate(headers) if h})
    return records


# ---------------------------------------------------------------------------
# Manifest traversal
# ---------------------------------------------------------------------------

def file_entry(manifest: dict[str, Any], source_type: str) -> dict[str, Any] | None:
    for item in manifest.get("files", []):
        if item.get("source_type") == source_type and item.get("parse_status") == "parsed":
            return item
    return None


def file_entries(manifest: dict[str, Any], source_type: str) -> list[dict[str, Any]]:
    return [
        item for item in manifest.get("files", [])
        if item.get("source_type") == source_type and item.get("parse_status") == "parsed"
    ]


def sheet_entry(file_item: dict[str, Any], role: str) -> dict[str, Any] | None:
    for sheet in file_item.get("sheets", []):
        if sheet.get("detected_role") == role:
            return sheet
    return None


def sheet_entries(file_item: dict[str, Any], role: str) -> list[dict[str, Any]]:
    return [s for s in file_item.get("sheets", []) if s.get("detected_role") == role]


def load_role_records(manifest: dict[str, Any], source_type: str, role: str) -> list[dict[str, Any]]:
    source_folder = Path(manifest["metadata"]["source_folder"])
    records: list[dict[str, Any]] = []
    for file_item in file_entries(manifest, source_type):
        for sheet in sheet_entries(file_item, role):
            xlsx_path = source_folder / file_item["relative_path"]
            column_remap: dict[str, str] = sheet.get("column_remap", {})
            for record in read_records(xlsx_path, sheet["sheet_name"], int(sheet["header_row"])):
                if column_remap:
                    for old_key, new_key in column_remap.items():
                        if old_key in record:
                            record[new_key] = record.pop(old_key)
                record["__source_file"] = file_item["file_name"]
                record["__source_sheet"] = sheet["sheet_name"]
                records.append(record)
    return records


# ---------------------------------------------------------------------------
# Record selection helpers
# ---------------------------------------------------------------------------

def first_by_value(records: list[dict[str, Any]], field: str, expected: str) -> dict[str, Any]:
    for record in records:
        if str(record.get(field, "")).strip() == expected:
            return record
    return records[0] if records else {}


def best_by_value(records: list[dict[str, Any]], field: str, expected: str, score_field: str) -> dict[str, Any]:
    matched = [r for r in records if str(r.get(field, "")).strip() == expected]
    if not matched:
        return {}
    return top_record(matched, score_field) or matched[0]


def first_by_source_and_value(
    records: list[dict[str, Any]], source_file: str | None, field: str, expected: str
) -> dict[str, Any]:
    if source_file:
        for record in records:
            if record.get("__source_file") == source_file and str(record.get(field, "")).strip() == expected:
                return record
    return first_by_value(records, field, expected)


def sum_top(records: list[dict[str, Any]], field: str, limit: int) -> float | None:
    values = [to_float(r.get(field)) for r in records[:limit]]
    values = [v for v in values if v is not None]
    return sum(values) if values else None


def top_record(records: list[dict[str, Any]], field: str) -> dict[str, Any]:
    best: dict[str, Any] = {}
    best_value = float("-inf")
    for record in records:
        value = to_float(record.get(field))
        if value is not None and value > best_value:
            best = record
            best_value = value
    return best


def top_records(records: list[dict[str, Any]], field: str, limit: int) -> list[dict[str, Any]]:
    return sorted(records, key=lambda r: to_float(r.get(field)) or 0, reverse=True)[:limit]
