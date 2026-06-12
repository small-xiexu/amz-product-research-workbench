from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_FINAL_FILES = ("report.md", "summary.md", "dashboard.html", "data.xlsx")
REQUIRED_WORKFLOW_FILES = ("workflow_summary.md", "workflow_summary.json")
BASE_REQUIRED_SHEETS = (
    "数据来源说明",
    "市场结构",
    "Top100原始明细",
    "数据质量检查",
    "属性定义",
    "Top商品打标",
    "待确认标签",
    "属性分布",
    "属性交叉分析",
    "机会判断",
    "竞品选择逻辑",
    "竞品池",
    "决策检查",
    "风险矩阵",
    "状态卡",
)
VOC_REQUIRED_SHEETS = ("评论VOC", "VOC证据")
PROFIT_REQUIRED_SHEETS = ("利润参考结果", "利润成本拆分")
IP_COMPLIANCE_REQUIRED_SHEETS = ("知产合规复核", "知产初筛", "合规认证预判")
REPORT_REQUIRED_TERMS = {
    "状态卡": ("状态卡",),
    "下一步": ("下一步",),
    "待补项": ("待补项", "待补"),
    "数据来源说明": ("数据来源说明", "数据来源"),
}
FORMAL_REPORT_SECTION_TITLES = (
    "Executive Summary / 当前结论",
    "数据来源与口径",
    "候选方向与边界",
    "市场结构与数据质量",
    "关键词与需求信号",
    "产品属性分布与交叉分析",
    "竞品池与竞品选择逻辑",
    "评论 VOC 与真实痛点",
    "利润复核",
    "知产/合规/退货风险",
    "Go/Wait/No-Go 决策检查",
    "下一步动作与证据附录",
)
WAITING_TEMPLATE_TEXT = "待填写模板"
TOP100_EXPECTED_ROWS = 100
REQUIRED_SCORECARD_DIMENSIONS = (
    "市场规模",
    "竞争格局",
    "需求清晰度",
    "新品友好度",
    "利润可行性",
    "知产/合规/退货风险",
    "数据完整度",
)


@dataclass
class ValidationResult:
    workflow_dir: Path
    final_report_dir: Path
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def to_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "workflow_dir": str(self.workflow_dir),
            "final_report_dir": str(self.final_report_dir),
            "errors": self.errors,
            "warnings": self.warnings,
            "notes": self.notes,
        }


def validate_workflow_output(input_dir: Path | str) -> ValidationResult:
    workflow_dir, final_report_dir = resolve_output_dirs(Path(input_dir).expanduser().resolve())
    result = ValidationResult(workflow_dir=workflow_dir, final_report_dir=final_report_dir)

    _check_required_files(result, final_report_dir, REQUIRED_FINAL_FILES, "final_report")
    _check_required_files(result, workflow_dir, REQUIRED_WORKFLOW_FILES, "workflow")

    workflow_summary = _load_json(workflow_dir / "workflow_summary.json", result)
    workflow_summary_text = _read_text(workflow_dir / "workflow_summary.md")
    report_text = _read_text(final_report_dir / "report.md")

    sheet_names: set[str] = set()
    data_workbook = final_report_dir / "data.xlsx"
    if data_workbook.exists():
        sheet_names = _load_sheet_names(data_workbook, result)
        _check_required_sheets(result, sheet_names, BASE_REQUIRED_SHEETS, "基础交付")
        _check_top100_rows(data_workbook, result)
        _check_attribute_analysis(data_workbook, result)
        _check_competitor_selection_logic(data_workbook, result)
        _check_go_nogo_scorecard(data_workbook, workflow_summary, result)

    if _review_voc_enabled(workflow_dir, workflow_summary):
        _check_required_sheets(result, sheet_names, VOC_REQUIRED_SHEETS, "评论 VOC")
        if data_workbook.exists():
            _check_voc_evidence_chain(data_workbook, result)

    profit_review = _section(workflow_summary, "profit_review")
    if bool(profit_review.get("applied")):
        _check_required_sheets(result, sheet_names, PROFIT_REQUIRED_SHEETS, "利润复核")
    else:
        _check_waiting_template_status(result, profit_review, workflow_summary_text, "利润复核")

    ip_compliance = _section(workflow_summary, "ip_compliance_review")
    if bool(ip_compliance.get("applied")):
        _check_required_sheets(result, sheet_names, IP_COMPLIANCE_REQUIRED_SHEETS, "知产/合规初筛")
    else:
        _check_waiting_template_status(result, ip_compliance, workflow_summary_text, "知产/合规初筛")

    _check_report_terms(result, report_text)
    _check_report_sections(result, report_text)
    return result


def resolve_output_dirs(path: Path) -> tuple[Path, Path]:
    if (path / "final_report").is_dir():
        return path, path / "final_report"
    if (path / "report.md").exists() or (path / "data.xlsx").exists():
        return path.parent, path
    return path, path / "final_report"


def render_result(result: ValidationResult) -> str:
    status = "通过" if result.ok else "失败"
    lines = [
        f"正式交付校验：{status}",
        f"- workflow_dir: {result.workflow_dir}",
        f"- final_report_dir: {result.final_report_dir}",
    ]
    if result.errors:
        lines.append("")
        lines.append("错误：")
        lines.extend(f"- {item}" for item in result.errors)
    if result.warnings:
        lines.append("")
        lines.append("警告：")
        lines.extend(f"- {item}" for item in result.warnings)
    if result.notes:
        lines.append("")
        lines.append("检查记录：")
        lines.extend(f"- {item}" for item in result.notes)
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate formal product research workflow outputs.")
    parser.add_argument("workflow_output_dir", help="run_research_workflow.py generated output directory.")
    parser.add_argument("--json", action="store_true", help="Print machine-readable validation result.")
    args = parser.parse_args(argv)

    result = validate_workflow_output(args.workflow_output_dir)
    if args.json:
        print(json.dumps(result.to_dict(), ensure_ascii=False, indent=2))
    else:
        print(render_result(result), end="")
    return 0 if result.ok else 1


def _check_required_files(result: ValidationResult, base_dir: Path, filenames: tuple[str, ...], label: str) -> None:
    for filename in filenames:
        path = base_dir / filename
        if path.exists():
            result.notes.append(f"{label} 文件存在：{filename}")
        else:
            result.errors.append(f"{label} 缺少必需文件：{path}")


def _load_json(path: Path, result: ValidationResult) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        result.errors.append(f"workflow_summary.json 不是合法 JSON：{exc}")
        return {}
    if isinstance(data, dict):
        return data
    result.errors.append("workflow_summary.json 根节点必须是对象")
    return {}


def _read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def _load_sheet_names(path: Path, result: ValidationResult) -> set[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl error types vary by file corruption mode
        result.errors.append(f"data.xlsx 无法读取：{exc}")
        return set()
    try:
        sheet_names = set(workbook.sheetnames)
        result.notes.append(f"data.xlsx 可读取，Sheet 数：{len(sheet_names)}")
        return sheet_names
    finally:
        workbook.close()


def _check_required_sheets(
    result: ValidationResult,
    sheet_names: set[str],
    required_sheets: tuple[str, ...],
    label: str,
) -> None:
    if not sheet_names:
        return
    missing = [name for name in required_sheets if name not in sheet_names]
    if missing:
        result.errors.append(f"{label} 缺少关键 Sheet：{', '.join(missing)}")
    else:
        result.notes.append(f"{label} 关键 Sheet 完整")


def _check_top100_rows(path: Path, result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        if "Top100原始明细" not in workbook.sheetnames:
            return
        row_count = _count_data_rows(workbook["Top100原始明细"])
    finally:
        workbook.close()
    if row_count < TOP100_EXPECTED_ROWS:
        result.warnings.append(f"Top100原始明细有效商品行数为 {row_count}，低于正式深挖建议的 {TOP100_EXPECTED_ROWS} 行")
    else:
        result.notes.append(f"Top100原始明细有效商品行数：{row_count}")


def _check_attribute_analysis(path: Path, result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        definitions = _count_data_rows(workbook["属性定义"]) if "属性定义" in workbook.sheetnames else 0
        cross_count = _count_group_rows(workbook["属性交叉分析"]) if "属性交叉分析" in workbook.sheetnames else 0
        opportunity_count = _count_data_rows(workbook["机会判断"]) if "机会判断" in workbook.sheetnames else 0
    finally:
        workbook.close()
    if definitions < 3:
        result.errors.append(f"属性定义少于 3 个核心维度：当前 {definitions} 个")
    else:
        result.notes.append(f"属性定义核心维度数：{definitions}")
    if cross_count < 3:
        result.errors.append(f"属性交叉分析少于 3 组：当前 {cross_count} 组")
    else:
        result.notes.append(f"属性交叉分析组数：{cross_count}")
    if opportunity_count < 1:
        result.errors.append("机会判断 Sheet 缺少真机会/伪机会/待验证记录")
    else:
        result.notes.append(f"机会判断记录数：{opportunity_count}")


def _check_competitor_selection_logic(path: Path, result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        if "竞品选择逻辑" not in workbook.sheetnames:
            return
        rows = list(workbook["竞品选择逻辑"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) <= 1:
        result.errors.append("竞品选择逻辑 Sheet 为空")
        return
    header = [str(item or "") for item in rows[0]]
    required = ["ASIN", "品牌", "价格", "月销量", "评分", "评分数", "竞品类型", "覆盖维度", "选择理由"]
    missing_headers = [item for item in required if not any(item in cell for cell in header)]
    if missing_headers:
        result.errors.append(f"竞品选择逻辑缺少字段：{', '.join(missing_headers)}")
        return
    data_rows = [row for row in rows[1:] if row and str(row[0] or "").strip() not in {"", "未生成"}]
    if not data_rows:
        result.errors.append("竞品选择逻辑缺少有效竞品行")
        return
    role_values = {str(row[7] or "") for row in data_rows if len(row) > 7}
    if len(role_values) < 3:
        result.warnings.append(f"竞品选择逻辑覆盖角色少于 3 类：{', '.join(sorted(role_values)) or '无'}")
    result.notes.append(f"竞品选择逻辑有效行数：{len(data_rows)}")


def _check_voc_evidence_chain(path: Path, result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        if "VOC证据" not in workbook.sheetnames:
            return
        rows = list(workbook["VOC证据"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) <= 1:
        result.errors.append("VOC证据 Sheet 为空，评论接入后必须保留 review_id、ASIN、评分和片段")
        return
    header = [str(item or "") for item in rows[0]]
    required = ["评论ID", "ASIN", "评分", "证据片段"]
    missing_headers = [item for item in required if item not in header]
    if missing_headers:
        result.errors.append(f"VOC证据缺少追溯字段：{', '.join(missing_headers)}")
        return
    indexes = {name: header.index(name) for name in required}
    valid_rows = 0
    incomplete_rows = 0
    for row in rows[1:]:
        if not row or str(row[0] or "").strip() == "未接入":
            continue
        values = [row[indexes[name]] if indexes[name] < len(row) else None for name in required]
        if all(str(value or "").strip() for value in values):
            valid_rows += 1
        else:
            incomplete_rows += 1
    if valid_rows == 0:
        result.errors.append("VOC证据缺少完整可追溯行：review_id、ASIN、评分、片段必须同时存在")
    if incomplete_rows:
        result.warnings.append(f"VOC证据存在 {incomplete_rows} 行追溯字段不完整")
    result.notes.append(f"VOC证据完整追溯行数：{valid_rows}")


def _check_go_nogo_scorecard(path: Path, workflow_summary: dict[str, Any], result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    try:
        if "Go_No-Go评分卡" not in workbook.sheetnames:
            return
        rows = list(workbook["Go_No-Go评分卡"].iter_rows(values_only=True))
    finally:
        workbook.close()
    dimension_names = {str(row[0] or "") for row in rows[1:] if row and str(row[0] or "") in REQUIRED_SCORECARD_DIMENSIONS}
    missing = [name for name in REQUIRED_SCORECARD_DIMENSIONS if name not in dimension_names]
    if missing:
        result.errors.append(f"Go/Wait/No-Go 评分卡缺少固定维度：{', '.join(missing)}")
    else:
        result.notes.append("Go/Wait/No-Go 评分卡 7 个固定维度完整")
    decision = ""
    gating_text = ""
    for row in rows:
        if row and row[0] == "决策结论":
            decision = str(row[1] or "")
        if row and row[0] == "决策限制":
            gating_text = str(row[1] or "")
    profit_applied = bool(_section(workflow_summary, "profit_review").get("applied"))
    ip_applied = bool(_section(workflow_summary, "ip_compliance_review").get("applied"))
    if (not profit_applied or not ip_applied) and decision.upper() == "GO":
        result.errors.append("利润或知产/合规未回填时，Go/Wait/No-Go 评分卡禁止输出 GO")
    if (not profit_applied or not ip_applied) and not gating_text:
        result.errors.append("利润或知产/合规未回填时，Go/Wait/No-Go 评分卡必须写明决策限制")


def _count_data_rows(sheet: Any) -> int:
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = [str(cell).strip() for cell in row if cell not in (None, "")]
        if not values:
            continue
        if all(value in {"待填", "未生成", "未接入"} for value in values):
            continue
        count += 1
    return count


def _count_group_rows(sheet: Any) -> int:
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        first_cell = str(row[0]).strip() if row and row[0] not in (None, "") else ""
        if first_cell and first_cell not in {"待填", "未生成", "未接入"}:
            count += 1
    return count


def _review_voc_enabled(workflow_dir: Path, workflow_summary: dict[str, Any]) -> bool:
    review_voc = _section(workflow_summary, "review_voc")
    return bool(review_voc.get("enabled")) or (workflow_dir / "review_voc_package.json").exists()


def _section(data: dict[str, Any], key: str) -> dict[str, Any]:
    value = data.get(key)
    return value if isinstance(value, dict) else {}


def _check_waiting_template_status(
    result: ValidationResult,
    section: dict[str, Any],
    workflow_summary_text: str,
    label: str,
) -> None:
    status = str(section.get("status", ""))
    if WAITING_TEMPLATE_TEXT in status or WAITING_TEMPLATE_TEXT in workflow_summary_text:
        result.notes.append(f"{label} 未回填状态已明确：{WAITING_TEMPLATE_TEXT}")
    else:
        result.errors.append(f"{label} 未回填时必须在 workflow_summary 中明确“{WAITING_TEMPLATE_TEXT}”")


def _check_report_terms(result: ValidationResult, report_text: str) -> None:
    if not report_text:
        return
    missing = [label for label, terms in REPORT_REQUIRED_TERMS.items() if not any(term in report_text for term in terms)]
    if missing:
        result.errors.append(f"report.md 缺少正式交付锚点：{', '.join(missing)}")
    else:
        result.notes.append("report.md 已包含状态卡、下一步、待补项和数据来源说明")


def _check_report_sections(result: ValidationResult, report_text: str) -> None:
    if not report_text:
        return
    positions: list[int] = []
    missing: list[str] = []
    for title in FORMAL_REPORT_SECTION_TITLES:
        marker = f"## {title}"
        position = report_text.find(marker)
        if position < 0:
            missing.append(title)
        else:
            positions.append(position)
    if missing:
        result.errors.append(f"report.md 缺少正式章节：{', '.join(missing)}")
        return
    if positions != sorted(positions):
        result.errors.append("report.md 正式章节顺序不符合 P20.2 标准")
        return
    result.notes.append("report.md 正式 12 章结构完整且顺序正确")


if __name__ == "__main__":
    raise SystemExit(main())
