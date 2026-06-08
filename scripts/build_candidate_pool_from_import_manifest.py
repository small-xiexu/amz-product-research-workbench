#!/usr/bin/env python3
"""Build a first candidate_pool from inspected manual export files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from packages.research_core.rules.market_structure_rules import build_market_structure_analysis

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit("openpyxl is required to read Excel exports") from exc


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def to_float(value: Any) -> float | None:
    if value in (None, "", "--", "--,--"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("$", "").replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def to_int(value: Any) -> int | None:
    number = to_float(value)
    return int(number) if number is not None else None


def slugify(text: str) -> str:
    lowered = text.lower()
    slug = re.sub(r"[^a-z0-9]+", "-", lowered).strip("-")
    return slug or "manual-export-market"


def read_records(xlsx_path: Path, sheet_name: str, header_row: int) -> list[dict[str, Any]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    header_values = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    headers = [str(clean_cell(value)) if clean_cell(value) not in (None, "") else "" for value in header_values]
    records: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = [clean_cell(value) for value in row]
        if not any(value not in (None, "") for value in values):
            continue
        record = {header: values[index] for index, header in enumerate(headers) if header}
        records.append(record)
    return records


def file_entry(manifest: dict[str, Any], source_type: str) -> dict[str, Any] | None:
    for item in manifest.get("files", []):
        if item.get("source_type") == source_type and item.get("parse_status") == "parsed":
            return item
    return None


def sheet_entry(file_item: dict[str, Any], role: str) -> dict[str, Any] | None:
    for sheet in file_item.get("sheets", []):
        if sheet.get("detected_role") == role:
            return sheet
    return None


def load_role_records(manifest: dict[str, Any], source_type: str, role: str) -> list[dict[str, Any]]:
    file_item = file_entry(manifest, source_type)
    if not file_item:
        return []
    sheet = sheet_entry(file_item, role)
    if not sheet:
        return []
    source_folder = Path(manifest["metadata"]["source_folder"])
    xlsx_path = source_folder / file_item["relative_path"]
    return read_records(xlsx_path, sheet["sheet_name"], int(sheet["header_row"]))


def first_by_value(records: list[dict[str, Any]], field: str, expected: str) -> dict[str, Any]:
    for record in records:
        if str(record.get(field, "")).strip() == expected:
            return record
    return records[0] if records else {}


def sum_top(records: list[dict[str, Any]], field: str, limit: int) -> float | None:
    values = [to_float(record.get(field)) for record in records[:limit]]
    values = [value for value in values if value is not None]
    if not values:
        return None
    return sum(values)


def top_record(records: list[dict[str, Any]], field: str) -> dict[str, Any]:
    best: dict[str, Any] = {}
    best_value = float("-inf")
    for record in records:
        value = to_float(record.get(field))
        if value is not None and value > best_value:
            best = record
            best_value = value
    return best


def product_from_concentration(record: dict[str, Any], note: str = "") -> dict[str, Any]:
    return {
        "asin": record.get("ASIN"),
        "title": record.get("商品标题"),
        "brand": record.get("品牌"),
        "seller": record.get("卖家"),
        "fulfillment": record.get("配送方式"),
        "price": to_float(record.get("价格($)")),
        "monthly_units": to_float(record.get("月销量")),
        "monthly_revenue_usd": to_float(record.get("月销售额($)")),
        "units_share": to_float(record.get("月销量占比")),
        "revenue_share": to_float(record.get("月销售额占比")),
        "rating": to_float(record.get("星级")),
        "rating_count": to_int(record.get("评分数")),
        "review_count": to_int(record.get("评论数")),
        "listing_date": record.get("上架时间"),
        "note": note,
    }


def product_from_search_result(record: dict[str, Any], note: str = "") -> dict[str, Any]:
    return {
        "asin": record.get("ASIN"),
        "title": record.get("商品标题"),
        "brand": record.get("品牌"),
        "seller": record.get("Buybox卖家"),
        "seller_location": record.get("卖家所属地"),
        "fulfillment": record.get("配送方式"),
        "price": to_float(record.get("价格($)")),
        "monthly_units": to_float(record.get("月销量")),
        "monthly_revenue_usd": to_float(record.get("月销售额($)")),
        "rating": to_float(record.get("评分")),
        "rating_count": to_int(record.get("评分数")),
        "bsr": to_int(record.get("小类BSR")),
        "category": record.get("小类目"),
        "listing_date": record.get("上架时间"),
        "listing_days": to_int(record.get("上架天数")),
        "parent_asin": record.get("父ASIN"),
        "variant_count": to_int(record.get("变体数")),
        "lqs": to_int(record.get("LQS")),
        "has_a_plus": record.get("A+页面"),
        "has_video": record.get("视频介绍"),
        "weight": record.get("商品重量（单位换算）") or record.get("商品重量"),
        "size": record.get("商品尺寸（单位换算）") or record.get("商品尺寸"),
        "package_weight": record.get("包装重量（单位换算）") or record.get("包装重量"),
        "package_size": record.get("包装尺寸（单位换算）") or record.get("包装尺寸"),
        "url": record.get("商品详情页链接"),
        "note": note,
    }


def top_products(product_concentration: list[dict[str, Any]], limit: int = 10) -> list[dict[str, Any]]:
    return [
        product_from_concentration(record, "Top 商品集中度")
        for record in product_concentration[:limit]
        if record.get("ASIN")
    ]


def top_products_from_search_results(search_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    seen: set[str] = set()
    for record in search_records:
        asin = str(record.get("ASIN") or "").strip()
        if not asin or asin in seen:
            continue
        seen.add(asin)
        result.append(product_from_search_result(record, "卖家精灵搜索结果明细"))
    return result


def recent_winners(search_records: list[dict[str, Any]], limit: int = 10) -> list[dict[str, Any]]:
    recent = [
        record
        for record in search_records
        if record.get("ASIN") and (to_int(record.get("上架天数")) or 999999) <= 180 and (to_float(record.get("月销量")) or 0) > 0
    ]
    recent.sort(key=lambda item: (to_float(item.get("月销量")) or 0, to_float(item.get("月销售额($)")) or 0), reverse=True)
    result = []
    seen: set[str] = set()
    for record in recent:
        asin = str(record.get("ASIN") or "")
        if not asin or asin in seen:
            continue
        seen.add(asin)
        result.append(product_from_search_result(record, "近半年上架且有销量"))
        if len(result) >= limit:
            break
    return result


def structure_supplements(search_records: list[dict[str, Any]], limit: int = 8) -> list[dict[str, Any]]:
    candidates: list[tuple[str, dict[str, Any]]] = []
    priced = [record for record in search_records if record.get("ASIN") and to_float(record.get("价格($)")) is not None]
    rated = [record for record in search_records if record.get("ASIN") and to_float(record.get("评分")) is not None]
    reviewed = [record for record in search_records if record.get("ASIN") and to_int(record.get("评分数")) is not None]
    if priced:
        candidates.append(("低价结构样本", min(priced, key=lambda item: to_float(item.get("价格($)")) or 0)))
        candidates.append(("高价结构样本", max(priced, key=lambda item: to_float(item.get("价格($)")) or 0)))
    if rated:
        candidates.append(("高评分结构样本", max(rated, key=lambda item: to_float(item.get("评分")) or 0)))
        low_rated = [record for record in rated if (to_float(record.get("评分")) or 0) < 4]
        if low_rated:
            candidates.append(("低评分风险样本", min(low_rated, key=lambda item: to_float(item.get("评分")) or 0)))
    if reviewed:
        candidates.append(("高评论门槛样本", max(reviewed, key=lambda item: to_int(item.get("评分数")) or 0)))

    seen: set[str] = set()
    result = []
    for note, record in candidates:
        asin = str(record.get("ASIN") or "")
        if not asin or asin in seen:
            continue
        seen.add(asin)
        result.append(product_from_search_result(record, note))
        if len(result) >= limit:
            break
    return result


def extract_seed_keyword(manifest: dict[str, Any]) -> str:
    search_file = file_entry(manifest, "seller_sprite_search_results")
    if search_file:
        match = re.search(r"Search\\((.*?)\\)", search_file.get("file_name", ""))
        if match:
            return match.group(1).replace("-", " ")
    aba_records = load_role_records(manifest, "amazon_aba_keywords", "aba_keywords")
    if aba_records:
        return str(aba_records[0].get("搜索词", "")).strip()
    return manifest.get("metadata", {}).get("task_name", "")


def build_candidate(manifest: dict[str, Any]) -> dict[str, Any]:
    overview_records = load_role_records(manifest, "seller_sprite_market_analysis", "market_overview")
    product_concentration = load_role_records(manifest, "seller_sprite_market_analysis", "product_concentration")
    brand_concentration = load_role_records(manifest, "seller_sprite_market_analysis", "brand_concentration")
    seller_location = load_role_records(manifest, "seller_sprite_market_analysis", "seller_location_distribution")
    demand_signal = load_role_records(manifest, "seller_sprite_market_analysis", "market_demand_signal")
    price_distribution = load_role_records(manifest, "seller_sprite_market_analysis", "price_distribution")
    listing_age = load_role_records(manifest, "seller_sprite_market_analysis", "listing_age_distribution")
    search_records = load_role_records(manifest, "seller_sprite_search_results", "product_candidates")
    reverse_keywords = load_role_records(manifest, "seller_sprite_reverse_asin_keywords", "reverse_asin_keywords")
    aba_keywords = load_role_records(manifest, "amazon_aba_keywords", "aba_keywords")

    all_products = first_by_value(overview_records, "样品分类", "全部商品")
    top10 = first_by_value(overview_records, "样品分类", "前10商品")
    new_products = first_by_value(overview_records, "样品分类", "6个月内上架")
    demand_12m = first_by_value(demand_signal, "范围", "12个月")
    top_brand = brand_concentration[0] if brand_concentration else {}
    top_product = product_concentration[0] if product_concentration else {}
    top_location = top_record(seller_location, "月销量") if seller_location else {}
    top_price_band = top_record(price_distribution, "月销量") if price_distribution else {}
    recent_listing = first_by_value(listing_age, "上架时间", "半年") if listing_age else {}
    top_keyword = top_record(reverse_keywords, "月搜索量") if reverse_keywords else {}
    top_aba = aba_keywords[0] if aba_keywords else {}
    seed_keyword = extract_seed_keyword(manifest)
    market_name = "Hands Free Leashes"
    candidate_id = "cand-" + slugify(seed_keyword or market_name)
    market_return_rate = to_float(demand_12m.get("市场退货率"))
    category_return_rate = to_float(demand_12m.get("同类目退货率"))
    return_level = "待确认"
    if market_return_rate is not None and category_return_rate is not None:
        return_level = "中" if market_return_rate > category_return_rate else "低"
    top_product_rows = top_products_from_search_results(search_records)
    market_structure = build_market_structure_analysis(top_product_rows, expected_count=100)

    source_refs = [
        f"manual_export:{item['file_name']}"
        for item in manifest.get("files", [])
        if item.get("parse_status") == "parsed" and item.get("source_type") != "system_file"
    ]

    return {
        "candidate_id": candidate_id,
        "name": market_name,
        "candidate_type": "market_direction",
        "status": "继续看",
        "reason": "卖家精灵手动导出数据已覆盖市场分析、搜索结果、关键词反查和 ABA，具备进入候选池初筛的基础证据。",
        "appearance_reason": [
            f"搜索入口：{seed_keyword}" if seed_keyword else "来自卖家精灵手动导出样例",
            f"市场样本商品数：{to_int(all_products.get('样本商品数'))}" if all_products else "市场样本数据待补",
            f"6个月内新品数量：{to_int(new_products.get('样本商品数'))}" if new_products else "新品数据待补",
        ],
        "demand_evidence": {
            "market_avg_monthly_units": to_float(all_products.get("月均销量")),
            "market_avg_monthly_revenue_usd": to_float(all_products.get("月均销售额($)")),
            "market_avg_price_usd": to_float(all_products.get("平均价格($)")),
            "top_keyword": top_keyword.get("关键词"),
            "top_keyword_monthly_searches": to_int(top_keyword.get("月搜索量")),
            "aba_top_search_term": top_aba.get("搜索词"),
            "aba_top_clicked_asin": top_aba.get("点击量最高的商品 #1：ASIN"),
        },
        "competition_structure": {
            "sample_product_count": to_int(all_products.get("样本商品数")),
            "top10_avg_monthly_units": to_float(top10.get("月均销量")),
            "top10_avg_monthly_revenue_usd": to_float(top10.get("月均销售额($)")),
            "top_product_asin": top_product.get("ASIN"),
            "top_product_monthly_units_share": to_float(top_product.get("月销量占比")),
            "top10_product_units_share": sum_top(product_concentration, "月销量占比", 10),
            "top_brand": top_brand.get("品牌"),
            "top_brand_units_share": to_float(top_brand.get("月销量占比")),
            "top_seller_location": top_location.get("卖家所属地"),
            "top_seller_location_units_share": to_float(top_location.get("销量占比")),
            "search_result_rows": len(search_records),
        },
        "competitor_candidates": {
            "top10": top_products(product_concentration),
            "recent_winners": recent_winners(search_records),
            "structure_supplement": structure_supplements(search_records),
        },
        "top_products": top_product_rows,
        "market_structure": market_structure,
        "new_listing_opportunity": {
            "new_listing_count_6m": to_int(new_products.get("样本商品数")),
            "new_listing_avg_monthly_units": to_float(new_products.get("月均销量")),
            "new_listing_avg_monthly_revenue_usd": to_float(new_products.get("月均销售额($)")),
            "recent_6m_units_share": to_float(recent_listing.get("销量占比")),
            "latest_listing_date": all_products.get("商品最新上架时间"),
        },
        "preliminary_profit_space": {
            "avg_price_usd": to_float(all_products.get("平均价格($)")),
            "top_price_band_by_units": top_price_band.get("价格区间($)"),
            "top_price_band_units_share": to_float(top_price_band.get("销量占比")),
            "note": "仅为价格空间参考，采购价、FBA、头程和入库配置费仍需后续利润复核。",
        },
        "risk_flags": [
            "市场退货率高于同类目平均" if return_level == "中" else "退货率暂未高于同类目平均",
            "宠物牵引绳涉及拉力、耐磨、扣具安全，需后续结合评论和合规/责任风险复核。",
        ],
        "return_risk": {
            "level": return_level,
            "market_return_rate": market_return_rate,
            "category_return_rate": category_return_rate,
            "source": "卖家精灵市场分析-商品需求趋势",
        },
        "ip_compliance_risk": {
            "level": "待确认",
            "notes": "需后续检查商标、外观/结构专利、宠物用品安全和材质宣称风险。",
        },
        "data_quality": {
            "source": "manual_export",
            "source_types": manifest.get("data_quality", {}).get("available_source_types", []),
            "missing_source_types": manifest.get("data_quality", {}).get("missing_source_types", []),
            "top_product_quality": market_structure.get("data_quality", {}),
        },
        "missing_data": [
            "采购价",
            "FBA费用",
            "头程费用",
            "入库配置费",
            "评论/VOC证据",
            "商标/专利复核",
        ],
        "next_step": "基于商品集中度表提取 Top10 标杆组、近半年新品组和结构补充组，再进入重点候选深挖。",
        "source_refs": source_refs,
    }


def build_candidate_pool(manifest: dict[str, Any]) -> dict[str, Any]:
    candidate = build_candidate(manifest)
    return {
        "metadata": {
            "pool_id": "pool-manual-export-" + datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S"),
            "site": manifest.get("metadata", {}).get("site") or "US",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "discovery_mode": "manual_export",
            "data_sources": manifest.get("data_quality", {}).get("available_source_types", []),
            "import_manifest": manifest.get("metadata", {}).get("manifest_id"),
        },
        "source_brief": {
            "brief_id": "brief-from-manual-export",
            "site": manifest.get("metadata", {}).get("site") or "US",
            "search_scope": {
                "free_text": manifest.get("metadata", {}).get("task_name", ""),
                "seed_keyword": extract_seed_keyword(manifest),
            },
            "exclusion_rules": [],
            "preference_rules": {},
        },
        "summary": {
            "total_candidates": 1,
            "continue_count": 1 if candidate["status"] == "继续看" else 0,
            "trial_count": 1 if candidate["status"] == "试做" else 0,
            "watch_count": 1 if candidate["status"] == "观察" else 0,
            "drop_count": 1 if candidate["status"] == "先放弃" else 0,
            "key_gaps": candidate["missing_data"],
        },
        "candidates": [candidate],
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build candidate_pool from manual export import manifest.")
    parser.add_argument("manifest", help="Path to import_manifest.json generated by inspect_manual_exports.py.")
    parser.add_argument("output", help="Path to write candidate_pool.json.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    manifest_path = Path(args.manifest).expanduser().resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    candidate_pool = build_candidate_pool(manifest)
    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(candidate_pool, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote candidate pool: {output}")
    print("Candidates:", len(candidate_pool["candidates"]))


if __name__ == "__main__":
    main()
