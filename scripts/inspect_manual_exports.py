#!/usr/bin/env python3
"""Inspect manual SellerSprite/Amazon exports and build an import manifest."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit("openpyxl is required to inspect Excel exports") from exc


REQUIRED_SOURCE_TYPES = [
    "seller_sprite_search_results",
    "seller_sprite_market_analysis",
    "seller_sprite_reverse_asin_keywords",
]

SHEET_RULES: dict[str, dict[str, list[str]]] = {
    "product_candidates": {
        "required": ["ASIN", "商品标题", "品牌", "类目路径", "月销量", "月销售额($)", "价格($)", "评分数", "评分", "上架时间"],
        "optional": ["父ASIN", "小类BSR", "FBA($)", "毛利率", "卖家所属地", "商品重量", "商品尺寸", "包装重量", "包装尺寸"],
    },
    "brand_summary": {
        "required": ["品牌", "月销量", "月销售额($)", "平均价格($)", "市场份额"],
        "optional": ["近1年销量", "近一年销售额($)"],
    },
    "seller_summary": {
        "required": ["卖家", "月销量", "月销售额($)", "平均价格($)", "市场份额"],
        "optional": ["近1年销量", "近一年销售额($)"],
    },
    "reverse_asin_keywords": {
        "required": ["关键词", "流量占比", "预估周曝光量", "自然排名", "月搜索量"],
        "optional": ["关键词翻译", "广告排名", "ABA周排名", "PPC价格", "建议竞价范围", "前十ASIN"],
    },
    "unique_words": {
        "required": ["词语", "出现频次", "百分比"],
        "optional": [],
    },
    "aba_keywords": {
        "required": ["搜索频率排名", "搜索词", "点击量最高的商品 #1：ASIN", "点击量最高的商品 #1：点击份额", "点击量最高的商品 #1：转化份额"],
        "optional": ["点击量最高的品牌 #1", "点击量最高的类别 #1", "报告日期"],
    },
    "market_overview": {
        "required": ["样品分类", "样本商品数", "月均销量", "月均销售额($)", "平均价格($)", "平均评分数", "平均星级"],
        "optional": ["商品首次上架时间", "商品最新上架时间"],
    },
    "market_keyword_trend": {
        "required": ["月份"],
        "optional": [],
    },
    "market_sales_trend": {
        "required": ["月份", "月销量", "月销售额($)", "平均BSR"],
        "optional": [],
    },
    "product_concentration": {
        "required": ["排名", "ASIN", "品牌", "商品标题", "月销量", "月销量占比", "月销售额($)", "月销售额占比"],
        "optional": ["卖家", "配送方式", "价格($)", "上架时间", "评分数", "星级", "评论数"],
    },
    "brand_concentration": {
        "required": ["排名", "品牌", "商品数量", "品牌月销量", "月销量占比", "品牌月销售额($)", "月销售额占比"],
        "optional": ["新品数量", "均价($)", "评分数", "星级", "评论数"],
    },
    "seller_concentration": {
        "required": ["排名", "卖家名称", "商品数量", "卖家月销量", "月销量占比", "卖家月销售额($)", "月销售额占比"],
        "optional": ["新品数量"],
    },
    "seller_type_distribution": {
        "required": ["卖家类型", "ASIN数量", "ASIN数量占比", "月销量", "月销量占比"],
        "optional": ["评分数", "评分值"],
    },
    "a_plus_video_distribution": {
        "required": ["A+视频类型", "产品数量", "产品数量占比", "产品销量", "产品销量占比"],
        "optional": [],
    },
    "seller_location_distribution": {
        "required": ["卖家所属地", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": [],
    },
    "market_demand_signal": {
        "required": ["范围", "市场退货率", "同类目退货率"],
        "optional": ["商品总数", "市场搜索购买比", "同类目搜索购买比"],
    },
    "listing_age_distribution": {
        "required": ["上架时间", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": [],
    },
    "listing_year_distribution": {
        "required": ["上架年份", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": [],
    },
    "rating_count_distribution": {
        "required": ["评分数区间", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": [],
    },
    "rating_value_distribution": {
        "required": ["评分值", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": ["评论总数"],
    },
    "price_distribution": {
        "required": ["价格区间($)", "产品数量", "月销量", "销量占比", "月销售额($)", "月销售额占比"],
        "optional": [],
    },
}

SHEET_NAME_ROLES = {
    "Brands": "brand_summary",
    "Sellers": "seller_summary",
    "Unique Words": "unique_words",
    "市场分析概况": "market_overview",
    "行业需求及趋势": "market_keyword_trend",
    "行业销售趋势": "market_sales_trend",
    "商品集中度": "product_concentration",
    "品牌集中度": "brand_concentration",
    "卖家集中度": "seller_concentration",
    "卖家类型分布": "seller_type_distribution",
    "A+视频分布": "a_plus_video_distribution",
    "卖家所属地分布": "seller_location_distribution",
    "商品需求趋势": "market_demand_signal",
    "上架时间分布": "listing_age_distribution",
    "上架趋势分布": "listing_year_distribution",
    "评分数分布": "rating_count_distribution",
    "评分值分布": "rating_value_distribution",
    "价格分布": "price_distribution",
}


def non_empty_count(row: tuple[Any, ...]) -> int:
    return sum(1 for cell in row if cell not in (None, ""))


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def clean_row(row: tuple[Any, ...]) -> list[Any]:
    return [clean_cell(cell) for cell in row]


def trim_empty_tail(values: list[Any]) -> list[Any]:
    end = len(values)
    while end and values[end - 1] in (None, ""):
        end -= 1
    return values[:end]


def detect_header(rows: list[tuple[Any, ...]]) -> tuple[int, list[str]]:
    if not rows:
        return 1, []
    candidates = [(idx, row) for idx, row in enumerate(rows, start=1) if non_empty_count(row)]
    if not candidates:
        return 1, []
    idx, row = max(candidates, key=lambda item: non_empty_count(item[1]))
    headers = [str(value).strip() for value in trim_empty_tail(clean_row(row)) if value not in (None, "")]
    return idx, headers


def detect_sheet_role(sheet_name: str, headers: list[str]) -> str:
    if sheet_name in SHEET_NAME_ROLES:
        return SHEET_NAME_ROLES[sheet_name]
    header_set = set(headers)
    if {"ASIN", "商品标题", "月销量", "月销售额($)"}.issubset(header_set):
        return "product_candidates"
    if {"关键词", "流量占比", "月搜索量", "自然排名"}.issubset(header_set):
        return "reverse_asin_keywords"
    if {"搜索频率排名", "搜索词", "点击量最高的商品 #1：ASIN"}.issubset(header_set):
        return "aba_keywords"
    return "notes" if len(headers) <= 1 else "unknown"


def detect_file_type(sheet_roles: list[str], suffix: str, file_name: str) -> tuple[str, str]:
    if suffix == ".pdf":
        return "pdf_snapshot", "market_snapshot"
    if suffix not in {".xlsx", ".csv"}:
        return "system_file" if file_name == ".DS_Store" else "unknown", "ignored"
    role_set = set(sheet_roles)
    if "market_overview" in role_set and "product_concentration" in role_set:
        return "seller_sprite_market_analysis", "market_analysis"
    if "product_candidates" in role_set:
        return "seller_sprite_search_results", "product_candidates"
    if "reverse_asin_keywords" in role_set:
        return "seller_sprite_reverse_asin_keywords", "keyword_reverse"
    if "aba_keywords" in role_set:
        return "amazon_aba_keywords", "aba_keywords"
    return "unknown", "unknown"


def sheet_quality(role: str, headers: list[str]) -> tuple[list[str], list[str], list[str]]:
    rule = SHEET_RULES.get(role, {"required": [], "optional": []})
    required = rule["required"]
    optional = rule["optional"]
    header_set = set(headers)
    missing = [field for field in required if field not in header_set]
    return required, optional, missing


def inspect_xlsx(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    warnings: list[str] = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True))
        header_row, headers = detect_header(rows)
        role = detect_sheet_role(worksheet.title, headers)
        required, optional, missing = sheet_quality(role, headers)
        sample_rows = [trim_empty_tail(clean_row(row)) for row in rows[header_row : min(header_row + 3, len(rows))]]
        sheet = {
            "sheet_name": worksheet.title,
            "rows": int(worksheet.max_row or 0),
            "cols": int(worksheet.max_column or 0),
            "header_row": header_row,
            "detected_role": role,
            "headers": headers,
            "required_fields": required,
            "missing_required_fields": missing,
            "optional_fields": optional,
            "sample_rows": sample_rows,
        }
        if missing and role not in {"notes", "unknown"}:
            warnings.append(f"{worksheet.title} 缺少必需字段: {', '.join(missing)}")
        sheets.append(sheet)
    return sheets, warnings


def inspect_csv(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
        reader = csv.reader(handle, dialect)
        rows = [tuple(row) for _, row in zip(range(10), reader)]
    header_row, headers = detect_header(rows)
    role = detect_sheet_role(path.stem, headers)
    required, optional, missing = sheet_quality(role, headers)
    sample_rows = [trim_empty_tail(clean_row(row)) for row in rows[header_row : min(header_row + 3, len(rows))]]
    warning = [f"{path.name} 缺少必需字段: {', '.join(missing)}"] if missing and role not in {"notes", "unknown"} else []
    return [
        {
            "sheet_name": path.stem,
            "rows": len(rows),
            "cols": max((len(row) for row in rows), default=0),
            "header_row": header_row,
            "detected_role": role,
            "headers": headers,
            "required_fields": required,
            "missing_required_fields": missing,
            "optional_fields": optional,
            "sample_rows": sample_rows,
        }
    ], warning


def inspect_file(path: Path, base: Path, index: int) -> dict[str, Any]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    sheets: list[dict[str, Any]] = []
    parse_status = "parsed"
    if path.name == ".DS_Store":
        parse_status = "ignored"
    elif suffix == ".xlsx":
        try:
            sheets, warnings = inspect_xlsx(path)
        except Exception as exc:  # pragma: no cover - defensive IO handling
            parse_status = "error"
            warnings.append(str(exc))
    elif suffix == ".csv":
        try:
            sheets, warnings = inspect_csv(path)
        except Exception as exc:  # pragma: no cover - defensive IO handling
            parse_status = "error"
            warnings.append(str(exc))
    elif suffix == ".pdf":
        parse_status = "unsupported"
        warnings.append("PDF can be kept as a visual reference; V1 structured parsing starts from Excel/CSV.")
    else:
        parse_status = "ignored"
    source_type, data_role = detect_file_type([sheet["detected_role"] for sheet in sheets], suffix, path.name)
    return {
        "file_id": f"file_{index:03d}",
        "file_name": path.name,
        "relative_path": str(path.relative_to(base)),
        "source_type": source_type,
        "source_tool": infer_source_tool(source_type),
        "data_role": data_role,
        "parse_status": parse_status,
        "warnings": warnings,
        "sheets": sheets,
    }


def infer_source_tool(source_type: str) -> str:
    return {
        "seller_sprite_search_results": "卖家精灵搜索结果/产品调研导出",
        "seller_sprite_market_analysis": "卖家精灵市场分析导出",
        "seller_sprite_reverse_asin_keywords": "卖家精灵关键词反查导出",
        "amazon_aba_keywords": "Amazon Brand Analytics 导出",
        "pdf_snapshot": "浏览器导出 PDF",
        "system_file": "系统文件",
    }.get(source_type, "未知来源")


def build_manifest(folder: Path, task_name: str | None, site: str | None) -> dict[str, Any]:
    files = [
        inspect_file(path, folder, index)
        for index, path in enumerate(sorted(p for p in folder.iterdir() if p.is_file()), start=1)
    ]
    available = sorted({item["source_type"] for item in files if item["parse_status"] == "parsed"})
    missing = [source_type for source_type in REQUIRED_SOURCE_TYPES if source_type not in available]
    warnings: list[str] = []
    if missing:
        warnings.append("缺少 V1 建议数据源: " + ", ".join(missing))
    for item in files:
        warnings.extend(f"{item['file_name']}: {warning}" for warning in item["warnings"])
    return {
        "metadata": {
            "manifest_id": f"manifest-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}",
            "task_name": task_name or folder.name,
            "site": site or "",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "source_folder": str(folder),
            "notes": "Generated from local manual export samples. Raw export files stay outside git.",
        },
        "files": files,
        "data_quality": {
            "available_source_types": available,
            "missing_source_types": missing,
            "warnings": warnings,
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Inspect manual SellerSprite/Amazon export files.")
    parser.add_argument("folder", help="Folder containing exported xlsx/csv/pdf files.")
    parser.add_argument("output", help="Path to write import_manifest.json.")
    parser.add_argument("--task-name", default=None, help="Optional task name for manifest metadata.")
    parser.add_argument("--site", default=None, help="Optional Amazon marketplace, such as US.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    folder = Path(args.folder).expanduser().resolve()
    if not folder.exists() or not folder.is_dir():
        raise SystemExit(f"Export folder does not exist or is not a directory: {folder}")
    manifest = build_manifest(folder, args.task_name, args.site)
    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote import manifest: {output}")
    print("Available source types:", ", ".join(manifest["data_quality"]["available_source_types"]))
    if manifest["data_quality"]["missing_source_types"]:
        print("Missing source types:", ", ".join(manifest["data_quality"]["missing_source_types"]))


if __name__ == "__main__":
    main()
