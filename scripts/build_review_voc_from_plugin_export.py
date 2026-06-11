#!/usr/bin/env python3
"""Extract and normalize review data from Amazon review plugin exports.

This script is a pure data tool: it reads the plugin Excel/HTML exports,
normalizes field names and formats, computes basic statistics, and writes
a clean structured JSON for Claude to analyze.

All VOC analysis (pain points, highlights, improvement hypotheses) is done
by Claude reading the normalized data in conversation — not by this script.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import sys
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required to read review plugin Excel exports") from exc

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from packages.report_renderer.render_report import _write_xlsx


EXPECTED_SHEET = "评论数据"
HEADER_ALIASES = {
    "ASIN": "asin",
    "站点": "site",
    "评论ID": "review_id",
    "评论地区": "review_region",
    "原始评论日期": "raw_date",
    "评论日期": "review_date",
    "原始评星": "raw_rating",
    "评星": "rating",
    "情绪": "sentiment",
    "评论人": "author",
    "评论人主页": "author_profile_url",
    "是否验证购买": "verified",
    "Helpful数量": "helpful_count",
    "是否有买家实拍": "has_buyer_image",
    "图片数量": "image_count",
    "图片链接": "image_urls",
    "是否有视频": "has_video",
    "评论产品的属性": "variant",
    "颜色": "color",
    "尺寸": "size",
    "英文标题": "title_en",
    "英文评论": "body_en",
    "标题中文翻译": "title_zh",
    "评论中文翻译": "body_zh",
    "评论链接": "url",
}


# ── HTML AI 报告解析 ──────────────────────────────────────────────────────────

class MarkdownBodyParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_style_or_script = False
        self.markdown_depth = 0
        self.chips: list[str] = []
        self.body_parts: list[str] = []
        self._capture_chip = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = {key: value or "" for key, value in attrs}
        class_name = attrs_dict.get("class", "")
        if tag in {"style", "script"}:
            self.in_style_or_script = True
        if tag == "div" and "markdown-body" in class_name:
            self.markdown_depth += 1
        elif self.markdown_depth and tag in {"p", "h1", "h2", "h3", "h4", "li"}:
            self.body_parts.append("\n")
        if tag == "span" and self._inside_chips(attrs_dict, class_name):
            self._capture_chip = True

    def handle_endtag(self, tag: str) -> None:
        if tag in {"style", "script"}:
            self.in_style_or_script = False
        if tag == "div" and self.markdown_depth:
            self.markdown_depth -= 1
        if tag == "span":
            self._capture_chip = False

    def handle_data(self, data: str) -> None:
        text = compact_text(data)
        if not text or self.in_style_or_script:
            return
        if self._capture_chip:
            self.chips.append(text)
        if self.markdown_depth:
            self.body_parts.append(text)

    @staticmethod
    def _inside_chips(attrs: dict[str, str], class_name: str) -> bool:
        return "chips" in attrs.get("data-parent-class", "") or "chip" in class_name


# ── 工具函数 ──────────────────────────────────────────────────────────────────

def compact_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clean_date(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return compact_text(value)


def to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ".").strip())
    except ValueError:
        return None


def to_int(value: Any) -> int:
    number = to_float(value)
    return int(number) if number is not None else 0


def distribution_items(values: list[Any], limit: int = 12) -> list[dict[str, Any]]:
    counter = Counter(compact_text(v) for v in values if compact_text(v))
    return [{"name": name, "count": count} for name, count in counter.most_common(limit)]


def format_distribution(items: list[dict[str, Any]], limit: int = 5) -> str:
    if not items:
        return "未识别"
    return "；".join(f"{item.get('name')} {item.get('count')}" for item in items[:limit])


def rating_bucket(value: Any) -> str:
    rating = to_float(value)
    if rating is None:
        return "未识别"
    return f"{int(round(rating))}星"


def is_low_rating(review: dict[str, Any]) -> bool:
    rating = to_float(review.get("rating"))
    return rating is not None and rating <= 3


def has_media(review: dict[str, Any]) -> bool:
    return (
        review.get("has_buyer_image") == "是"
        or review.get("has_video") == "是"
        or to_int(review.get("image_count")) > 0
    )


def infer_sentiment(rating: float | None) -> str:
    if rating is None:
        return ""
    if rating >= 4:
        return "正面"
    if rating == 3:
        return "中性"
    if rating > 0:
        return "负面"
    return ""


# ── 数据读取与规范化 ──────────────────────────────────────────────────────────

def read_review_excel(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = EXPECTED_SHEET if EXPECTED_SHEET in workbook.sheetnames else workbook.sheetnames[0]
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return []
    headers = [HEADER_ALIASES.get(compact_text(v), compact_text(v)) for v in header_values]
    records: list[dict[str, Any]] = []
    for row in rows:
        if not any(v not in (None, "") for v in row):
            continue
        raw = {headers[i]: v for i, v in enumerate(row) if i < len(headers)}
        record = normalize_review(raw, path)
        if record.get("review_id") or record.get("review_text") or record.get("review_text_zh"):
            records.append(record)
    return records


def normalize_review(raw: dict[str, Any], source_path: Path) -> dict[str, Any]:
    rating = to_float(raw.get("rating"))
    title_en = compact_text(raw.get("title_en"))
    body_en = compact_text(raw.get("body_en"))
    title_zh = compact_text(raw.get("title_zh"))
    body_zh = compact_text(raw.get("body_zh"))
    review_text = compact_text(" ".join(p for p in [title_en, body_en] if p))
    review_text_zh = compact_text(" ".join(p for p in [title_zh, body_zh] if p))
    review_id = compact_text(raw.get("review_id")) or _fallback_review_id(raw)
    return {
        "review_id": review_id,
        "asin": compact_text(raw.get("asin")),
        "site": compact_text(raw.get("site")),
        "review_region": compact_text(raw.get("review_region")),
        "raw_date": compact_text(raw.get("raw_date")),
        "review_date": clean_date(raw.get("review_date")),
        "raw_rating": compact_text(raw.get("raw_rating")),
        "rating": rating if rating is not None else compact_text(raw.get("rating")),
        "sentiment": compact_text(raw.get("sentiment")) or infer_sentiment(rating),
        "author": compact_text(raw.get("author")),
        "verified": compact_text(raw.get("verified")),
        "helpful_count": to_int(raw.get("helpful_count")),
        "has_buyer_image": compact_text(raw.get("has_buyer_image")),
        "image_count": to_int(raw.get("image_count")),
        "has_video": compact_text(raw.get("has_video")),
        "variant": compact_text(raw.get("variant")),
        "color": compact_text(raw.get("color")),
        "size": compact_text(raw.get("size")),
        "review_text": review_text,
        "review_text_zh": review_text_zh,
        "url": compact_text(raw.get("url")),
        "source_file": source_path.name,
    }


def _fallback_review_id(raw: dict[str, Any]) -> str:
    source = "|".join(compact_text(raw.get(k)) for k in ("asin", "review_date", "author", "body_en", "body_zh"))
    return "review-" + hashlib.sha1(source.encode("utf-8")).hexdigest()[:12]


def read_ai_report(path: Path) -> dict[str, Any]:
    parser = MarkdownBodyParser()
    parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
    text = "\n".join(part for part in parser.body_parts if compact_text(part))
    return {
        "source_file": path.name,
        "text": compact_text(text),
    }


# ── 数据包构建 ────────────────────────────────────────────────────────────────

def build_voc_package(
    reviews: list[dict[str, Any]],
    ai_reports: list[dict[str, Any]],
    source_files: list[Path],
    candidate_id: str,
    candidate_name: str,
) -> dict[str, Any]:
    """Build a clean review data package for Claude to analyze.

    No rule-based VOC analysis is performed here. Claude reads
    normalized_reviews and stats to provide real analysis in conversation.
    """
    now = datetime.now(timezone.utc).isoformat()
    stats = _build_stats(reviews)
    return {
        "metadata": {
            "package_id": "review-voc-" + datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S"),
            "generated_at": now,
            "source_type": "review_plugin_export",
            "candidate_id": candidate_id,
            "candidate_name": candidate_name,
            "source_files": [str(p) for p in source_files],
            "integration_note": "Excel is the structured source of truth; HTML AI reports are auxiliary reading references.",
            "analysis_note": "痛点/亮点/改品机会分析由 Claude 在对话中基于 normalized_reviews 完成，脚本只做数据提取和规范化。",
        },
        "stats": stats,
        "ai_report_reference": {
            "source_files": [r["source_file"] for r in ai_reports],
            "text_excerpt": "\n\n".join(r["text"][:1800] for r in ai_reports if r.get("text")),
            "usage": "仅作辅助阅读，不作为可追溯证据主来源。",
        },
        "normalized_reviews": reviews,
        # Legacy fields kept empty so downstream code doesn't break.
        # Real analysis is done by Claude reading normalized_reviews.
        "summary": stats,
        "pain_points": [],
        "highlights": [],
        "opportunity_hypotheses": [],
    }


def _build_stats(reviews: list[dict[str, Any]]) -> dict[str, Any]:
    ratings = Counter(rating_bucket(r.get("rating")) for r in reviews)
    sentiments = Counter(compact_text(r.get("sentiment")) or "未识别" for r in reviews)
    asins = sorted({r.get("asin") for r in reviews if r.get("asin")})
    entry_site_dist = distribution_items([r.get("site") for r in reviews])
    region_dist = distribution_items([r.get("review_region") for r in reviews])
    dates = sorted(r.get("review_date") for r in reviews if r.get("review_date"))
    primary_entry_site = entry_site_dist[0]["name"] if entry_site_dist else ""
    primary_review_region = region_dist[0]["name"] if region_dist else ""
    return {
        "review_count": len(reviews),
        "asin_count": len(asins),
        "asins": asins,
        "date_range": {"start": dates[0] if dates else "", "end": dates[-1] if dates else ""},
        "rating_distribution": dict(sorted(ratings.items())),
        "sentiment_distribution": dict(sentiments),
        "low_rating_count": sum(1 for r in reviews if is_low_rating(r)),
        "media_review_count": sum(1 for r in reviews if has_media(r)),
        "entry_site_distribution": entry_site_dist,
        "review_region_distribution": region_dist,
        "primary_entry_site": primary_entry_site,
        "primary_review_region": primary_review_region,
        "source_scope_note": (
            "站点字段表示评论采集入口，不等同于目标市场；评论地区表示评论样本实际地区。"
            "跨站采集时，应结合目标站点、评论地区和评论内容共同解释。"
        ),
    }


# ── 输出渲染 ──────────────────────────────────────────────────────────────────

def render_markdown(package: dict[str, Any]) -> str:
    meta = package.get("metadata", {})
    stats = package.get("stats") or package.get("summary", {})
    lines = [
        f"# {meta.get('candidate_name') or '评论数据'} — 评论数据概况",
        "",
        "> 本文件只包含数据概况，不包含规则生成的痛点/亮点分析。",
        "> 请 Claude 读取 `normalized_reviews` 字段进行实质 VOC 分析。",
        "",
        "## 数据统计",
        f"- 总评论数：{stats.get('review_count', 0)}",
        f"- 覆盖 ASIN 数：{stats.get('asin_count', 0)}",
        f"- 低分评论数（≤3星）：{stats.get('low_rating_count', 0)}",
        f"- 含图片/视频评论数：{stats.get('media_review_count', 0)}",
        f"- 采集入口站点：{format_distribution(stats.get('entry_site_distribution', []))}",
        f"- 评论地区分布：{format_distribution(stats.get('review_region_distribution', []))}",
        f"- 时间范围：{stats.get('date_range', {}).get('start', '')} 至 {stats.get('date_range', {}).get('end', '')}",
        f"- 口径说明：{stats.get('source_scope_note', '')}",
        "",
        "## 评分分布",
    ]
    for star, count in sorted(stats.get("rating_distribution", {}).items()):
        lines.append(f"- {star}：{count} 条")
    lines += [
        "",
        "## 覆盖 ASIN",
    ]
    for asin in stats.get("asins", []):
        lines.append(f"- {asin}")
    lines += [
        "",
        "## AI 报告参考",
        f"- 来源：{', '.join(package.get('ai_report_reference', {}).get('source_files', []) or ['无'])}",
        "- 用途：仅作辅助阅读，正式结论必须追溯到评论 ID。",
    ]
    return "\n".join(lines) + "\n"


def render_summary(package: dict[str, Any]) -> str:
    stats = package.get("stats") or package.get("summary", {})
    return "\n".join([
        "# 评论数据摘要",
        "",
        f"- 评论数：{stats.get('review_count', 0)}",
        f"- 覆盖 ASIN：{stats.get('asin_count', 0)} 个",
        f"- 低分评论：{stats.get('low_rating_count', 0)} 条",
        f"- 采集入口：{format_distribution(stats.get('entry_site_distribution', []), 2)}",
        f"- 评论地区：{format_distribution(stats.get('review_region_distribution', []), 3)}",
        "",
        "> 痛点/亮点分析由 Claude 在对话中完成。",
    ]) + "\n"


def render_workbook(package: dict[str, Any], output_path: Path) -> None:
    stats = package.get("stats") or package.get("summary", {})
    _write_xlsx(
        output_path,
        [
            ("数据概况", _stats_rows(stats)),
            ("评论明细", _review_rows(package.get("normalized_reviews", []))),
            ("AI报告参考", _ai_report_rows(package.get("ai_report_reference", {}))),
        ],
    )


def _stats_rows(stats: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = [["字段", "值"]]
    simple_fields = [
        ("review_count", "总评论数"),
        ("asin_count", "覆盖ASIN数"),
        ("low_rating_count", "低分评论数(≤3星)"),
        ("media_review_count", "含图片/视频评论数"),
        ("primary_entry_site", "主要采集入口站点"),
        ("primary_review_region", "主要评论地区"),
        ("source_scope_note", "口径说明"),
    ]
    for key, label in simple_fields:
        rows.append([label, stats.get(key, "")])
    rows.append(["时间范围起", stats.get("date_range", {}).get("start", "")])
    rows.append(["时间范围止", stats.get("date_range", {}).get("end", "")])
    rows.append(["", ""])
    rows.append(["评分分布", ""])
    for star, count in sorted(stats.get("rating_distribution", {}).items()):
        rows.append([star, count])
    rows.append(["", ""])
    rows.append(["覆盖ASIN", ""])
    for asin in stats.get("asins", []):
        rows.append([asin, ""])
    return rows


def _review_rows(reviews: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = [["评论ID", "ASIN", "采集入口站点", "评论地区", "评分", "情绪", "日期", "Helpful", "属性", "中文评论", "英文评论", "链接"]]
    for r in reviews:
        rows.append([
            r.get("review_id"), r.get("asin"), r.get("site"), r.get("review_region"),
            r.get("rating"), r.get("sentiment"), r.get("review_date"),
            r.get("helpful_count"), r.get("variant"),
            r.get("review_text_zh"), r.get("review_text"), r.get("url"),
        ])
    return rows


def _ai_report_rows(reference: dict[str, Any]) -> list[list[Any]]:
    return [
        ["字段", "值"],
        ["来源文件", "\n".join(reference.get("source_files", []))],
        ["用途", reference.get("usage", "")],
        ["文本摘录", reference.get("text_excerpt", "")],
    ]


def display_value(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


# ── CLI ───────────────────────────────────────────────────────────────────────

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract and normalize review data from Amazon review plugin exports.")
    parser.add_argument("output_dir", help="Directory to write review_voc_package.json and report outputs.")
    parser.add_argument("inputs", nargs="+", help="Review plugin .xlsx exports and optional .html AI reports.")
    parser.add_argument("--candidate-id", default="", help="Candidate id from candidate_pool, if available.")
    parser.add_argument("--candidate-name", default="", help="Candidate name, if available.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_paths = [Path(item).expanduser().resolve() for item in args.inputs]
    excel_paths = [p for p in input_paths if p.suffix.lower() == ".xlsx"]
    html_paths = [p for p in input_paths if p.suffix.lower() in {".html", ".htm"}]
    if not excel_paths:
        raise SystemExit("At least one review plugin Excel export is required.")

    reviews: list[dict[str, Any]] = []
    for path in excel_paths:
        if not path.exists():
            raise SystemExit(f"Review input does not exist or is not a file: {path}")
        reviews.extend(read_review_excel(path))

    ai_reports = [read_ai_report(p) for p in html_paths]
    package = build_voc_package(reviews, ai_reports, input_paths, args.candidate_id, args.candidate_name)

    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "review_voc_package.json").write_text(
        json.dumps(package, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (output_dir / "voc_report.md").write_text(render_markdown(package), encoding="utf-8")
    (output_dir / "voc_summary.md").write_text(render_summary(package), encoding="utf-8")
    render_workbook(package, output_dir / "voc_evidence.xlsx")
    print(f"Wrote review data package: {output_dir / 'review_voc_package.json'}")
    print(f"Reviews: {len(reviews)}")
    print(f"ASINs: {package['stats']['asin_count']}")
    print(f"Low-rating reviews: {package['stats']['low_rating_count']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
