#!/usr/bin/env python3
"""Apply manual IP/compliance screening results to a research package."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit("openpyxl is required to read IP/compliance templates") from exc


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


RISK_SCORE = {
    "低": 1,
    "待复核": 2,
    "中": 3,
    "高": 4,
    "强风险": 5,
    "": 0,
}

RISK_ACTION = {
    "低": "保留检索记录，继续分析。",
    "中": "继续试做/观察，并保留待复核项。",
    "高": "暂停推进，先做人工或专业复核。",
    "强风险": "建议淘汰，或更换设计/功能/关键词后重新评估。",
    "待复核": "补充证据链接、截图或专业复核意见。",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply filled IP/compliance review template to research_package.json.")
    parser.add_argument("research_package_json")
    parser.add_argument("ip_compliance_template_xlsx")
    parser.add_argument("output_research_package_json")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    package_path = Path(args.research_package_json).expanduser().resolve()
    template_path = Path(args.ip_compliance_template_xlsx).expanduser().resolve()
    output_path = Path(args.output_research_package_json).expanduser().resolve()
    package = json.loads(package_path.read_text(encoding="utf-8"))
    review = read_ip_compliance_review(template_path)
    updated = apply_ip_compliance_review(package, review)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(updated, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote research package with IP/compliance review: {output_path}")
    return 0


def read_ip_compliance_review(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return {
        "product_flags": read_key_value_sheet(workbook, "产品属性", key_field="field", value_field="value"),
        "ip_rows": read_table_sheet(workbook, "知产初筛"),
        "compliance_rows": read_table_sheet(workbook, "合规初筛"),
    }


def apply_ip_compliance_review(package: dict[str, Any], review: dict[str, Any]) -> dict[str, Any]:
    result = json.loads(json.dumps(package, ensure_ascii=False))
    ip_summary = summarize_rows(review.get("ip_rows", []), "risk_type")
    compliance_summary = summarize_rows(review.get("compliance_rows", []), "product_attribute")
    overall_level = max_level([ip_summary["level"], compliance_summary["level"]])
    missing = missing_review_items(review)
    pending = pending_review_items(review)
    status = review_status(missing, pending)

    result["ip_screening"] = {
        "status": status,
        "level": ip_summary["level"],
        "overall_level": overall_level,
        "summary": ip_summary["summary"],
        "rows": review.get("ip_rows", []),
        "boundary": "仅为早期初筛，不替代律师或专利代理机构结论。",
    }
    result["compliance_screening"] = {
        "status": status,
        "level": compliance_summary["level"],
        "overall_level": overall_level,
        "product_flags": review.get("product_flags", {}),
        "summary": compliance_summary["summary"],
        "rows": review.get("compliance_rows", []),
        "boundary": "仅为可能材料和待复核项，不替代检测机构或 Amazon 合规团队结论。",
    }
    result["ip_compliance_review"] = {
        "status": status,
        "overall_level": overall_level,
        "ip_level": ip_summary["level"],
        "compliance_level": compliance_summary["level"],
        "ip_summary": ip_summary["summary"],
        "compliance_summary": compliance_summary["summary"],
        "missing_fields": missing,
        "pending_fields": pending,
        "next_step": next_step_for(overall_level, missing, pending),
    }
    result["decision_review"] = update_decision_review(result.get("decision_review", {}), result["ip_compliance_review"])
    result["status_card"] = update_status_card(result.get("status_card", {}), result["ip_compliance_review"])
    return result


def read_table_sheet(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        return []
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [compact_text(value) for value in rows[0]]
    records = []
    for row in rows[1:]:
        record = {headers[index]: clean_value(row[index] if index < len(row) else None) for index in range(len(headers)) if headers[index]}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def read_key_value_sheet(workbook: Any, sheet_name: str, key_field: str, value_field: str) -> dict[str, Any]:
    result = {}
    for record in read_table_sheet(workbook, sheet_name):
        key = compact_text(record.get(key_field))
        if key:
            result[key] = clean_value(record.get(value_field))
    return result


def summarize_rows(rows: list[dict[str, Any]], label_field: str) -> dict[str, Any]:
    reviewed = [row for row in rows if compact_text(row.get("result"))]
    if not reviewed:
        return {"level": "待复核", "summary": "尚未填写人工初筛结果。"}
    highest = max((normalize_level(row.get("result")) for row in reviewed), key=lambda level: RISK_SCORE.get(level, 0))
    grouped = group_review_rows(reviewed, label_field)
    risky = [item for item in grouped.values() if RISK_SCORE.get(item["level"], 0) >= RISK_SCORE["待复核"]]
    if risky:
        pieces = []
        for item in sorted(risky, key=lambda value: RISK_SCORE.get(value["level"], 0), reverse=True)[:5]:
            note = item["evidence_notes"][0] if item["evidence_notes"] else "见证据链接"
            pieces.append(f"{item['label']}：{item['level']}，{note}")
        summary = "；".join(pieces)
    else:
        summary = "已填写初筛结果，当前未记录中高风险项。"
    return {"level": highest, "summary": summary}


def group_review_rows(rows: list[dict[str, Any]], label_field: str) -> dict[str, dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        label = compact_text(row.get(label_field)) or "风险项"
        level = normalize_level(row.get("result"))
        note = compact_text(row.get("evidence_note"))
        item = grouped.setdefault(label, {"label": label, "level": level, "evidence_notes": []})
        current_score = RISK_SCORE.get(item["level"], 0)
        level_score = RISK_SCORE.get(level, 0)
        if level_score > current_score:
            item["level"] = level
            item["evidence_notes"] = []
        if level_score == RISK_SCORE.get(item["level"], 0) and note and note not in item["evidence_notes"]:
            item["evidence_notes"].append(note)
    return grouped


def missing_review_items(review: dict[str, Any]) -> list[str]:
    missing = []
    flags = review.get("product_flags", {})
    for key in ("product_usage", "user_group", "material_coating", "package_instruction_plan"):
        if not compact_text(flags.get(key)):
            missing.append(product_flag_label(key))
    if not any(compact_text(row.get("result")) for row in review.get("ip_rows", [])):
        missing.append("知产初筛结果")
    if not any(compact_text(row.get("result")) for row in review.get("compliance_rows", [])):
        missing.append("合规初筛结果")
    return missing


def pending_review_items(review: dict[str, Any]) -> list[str]:
    pending: list[str] = []
    for row in review.get("ip_rows", []):
        if normalize_level(row.get("result")) == "待复核":
            append_unique(pending, f"知产：{compact_text(row.get('risk_type')) or '未命名风险项'}")
    for row in review.get("compliance_rows", []):
        if normalize_level(row.get("result")) == "待复核":
            append_unique(pending, f"合规：{compact_text(row.get('product_attribute')) or '未命名规则'}")
    return pending


def append_unique(values: list[str], value: str) -> None:
    if value and value not in values:
        values.append(value)


def review_status(missing: list[str], pending: list[str]) -> str:
    if missing:
        return "待补"
    if pending:
        return "待复核"
    return "已初筛"


def product_flag_label(key: str) -> str:
    return {
        "product_usage": "产品用途",
        "user_group": "使用人群",
        "material_coating": "材质和涂层",
        "package_instruction_plan": "包装和说明书计划",
    }.get(key, key)


def update_decision_review(decision: dict[str, Any], review: dict[str, Any]) -> dict[str, Any]:
    updated = dict(decision or {})
    missing = set(review.get("missing_fields", []))
    pending = set(review.get("pending_fields", []))
    missing_inputs = [
        item
        for item in updated.get("missing_inputs", [])
        if item not in {"商标/专利复核", "合规认证复核", "知产初筛结果", "合规初筛结果"}
        and not str(item).startswith("知产：")
        and not str(item).startswith("合规：")
    ]
    for item in [*review.get("missing_fields", []), *review.get("pending_fields", [])]:
        if item not in missing_inputs:
            missing_inputs.append(item)
    facts = list(updated.get("facts", []))
    if not missing and not pending:
        facts.append(f"知产/合规初筛：整体风险 {review.get('overall_level')}")
    action_items = replace_ip_compliance_action(updated.get("action_items", []), review)
    updated["facts"] = facts
    updated["missing_inputs"] = missing_inputs
    updated["action_items"] = action_items
    updated["risk_matrix"] = update_risk_matrix(updated.get("risk_matrix", []), review)
    return updated


def replace_ip_compliance_action(action_items: list[str], review: dict[str, Any]) -> list[str]:
    # Action items are Claude's work. Return only non-IP filtered items.
    return [item for item in action_items if "商标" not in item and "专利" not in item and "合规" not in item and "知产" not in item]


def update_risk_matrix(risks: list[dict[str, Any]], review: dict[str, Any]) -> list[dict[str, Any]]:
    unresolved = [*review.get("missing_fields", []), *review.get("pending_fields", [])]
    summaries = [
        compact_text(review.get("ip_summary")),
        compact_text(review.get("compliance_summary")),
    ]
    summary_text = "；".join(item for item in summaries if item and "尚未填写" not in item) or "人工初筛已回填，详见知产初筛和合规认证预判。"
    basis = f"待补/待复核：{'、'.join(unresolved)}；{summary_text}" if unresolved else summary_text
    item = {
        "dimension": "知产/合规",
        "level": review.get("overall_level", "待复核"),
        "basis": basis,
        "next_check": review.get("next_step", "继续补充人工初筛结果。"),
    }
    updated = []
    replaced = False
    for risk in risks:
        if risk.get("dimension") == "知产/合规":
            updated.append(item)
            replaced = True
        else:
            updated.append(risk)
    if not replaced:
        updated.append(item)
    return updated


def update_status_card(status: dict[str, Any], review: dict[str, Any]) -> dict[str, Any]:
    updated = dict(status or {})
    updated["next_step"] = review.get("next_step", updated.get("next_step", "待补"))
    return updated


def next_step_for(level: str, missing: list[str], pending: list[str]) -> str:
    if missing:
        return "补齐知产/合规初筛字段：" + "、".join(missing)
    if pending:
        return "先处理待复核项：" + "、".join(pending)
    if level in {"强风险", "高"}:
        action = RISK_ACTION.get(level, "暂停推进，先做专业复核。")
        return f"{action} 复核完成前不要进入打样或备货。"
    if level == "中":
        return "保留当前候选，但在打样前补充证据链接、供应商资质和必要检测/专利复核。"
    if level == "低":
        return "保留检索记录，继续结合利润复核和供应链打样判断是否推进。"
    return "继续补充人工初筛结果，并保留证据链接和复核备注。"


def max_level(levels: list[str]) -> str:
    normalized = [normalize_level(level) for level in levels]
    return max(normalized, key=lambda level: RISK_SCORE.get(level, 0)) if normalized else "待复核"


def normalize_level(value: Any) -> str:
    text = compact_text(value)
    if text in RISK_SCORE:
        return text
    if "强" in text:
        return "强风险"
    if "高" in text:
        return "高"
    if "中" in text:
        return "中"
    if "低" in text:
        return "低"
    return "待复核"


def clean_value(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def compact_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


if __name__ == "__main__":
    raise SystemExit(main())
