#!/usr/bin/env python3
"""Apply a filled profit review template to a research package."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit("openpyxl is required to read profit templates") from exc


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from packages.research_core.rules.profit_rules import (
    DEFAULT_AD_RATE,
    DEFAULT_COMMISSION_RATE,
    DEFAULT_STORAGE_FEE_RATE,
    FIRST_LEG_RATE_CNY_PER_KG,
    ProfitInputs,
    calc_billable_weight_g,
    calc_first_leg_shipping,
    calc_profit_breakdown,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply filled profit review template to research_package.json.")
    parser.add_argument("research_package_json")
    parser.add_argument("profit_template_xlsx")
    parser.add_argument("output_research_package_json")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    package_path = Path(args.research_package_json).expanduser().resolve()
    template_path = Path(args.profit_template_xlsx).expanduser().resolve()
    output_path = Path(args.output_research_package_json).expanduser().resolve()

    package = json.loads(package_path.read_text(encoding="utf-8"))
    inputs = read_profit_inputs(template_path)
    updated = apply_profit_review(package, inputs)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(updated, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote research package with profit review: {output_path}")
    return 0


def read_profit_inputs(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "利润输入" not in workbook.sheetnames:
        raise SystemExit("profit template must contain sheet: 利润输入")
    worksheet = workbook["利润输入"]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result: dict[str, Any] = {}
    for row in rows[1:]:
        record = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
        field = compact_text(record.get("field"))
        if not field:
            continue
        result[field] = clean_value(record.get("value"))
    return result


def apply_profit_review(package: dict[str, Any], raw_inputs: dict[str, Any]) -> dict[str, Any]:
    result = json.loads(json.dumps(package, ensure_ascii=False))
    currency_code = site_currency_code(result.get("metadata", {}).get("site"))
    normalized = normalize_profit_inputs(raw_inputs, result)
    missing = validate_profit_inputs(normalized)
    result["operator_inputs"] = {
        **result.get("operator_inputs", {}),
        **normalized,
    }
    result["profit_reference"] = build_profit_reference(normalized, missing)
    result["profit_reference"]["currency_code"] = currency_code
    result["profit_review"] = {
        "status": "待补" if missing else "已计算",
        "missing_fields": missing,
        "input_source": "profit_template",
        "currency_code": currency_code,
        "currency_note": f"售价、FBA 和结果为 {currency_code}；采购价、头程、入库配置费按人民币输入后用汇率换算。",
    }
    result["decision_review"] = update_decision_review(result.get("decision_review", {}), missing, result["profit_reference"])
    result["status_card"] = update_status_card(result.get("status_card", {}), missing)
    return result


def normalize_profit_inputs(raw: dict[str, Any], package: dict[str, Any]) -> dict[str, Any]:
    exchange_rate = to_float(raw.get("exchange_rate"))
    sale_price = to_float(raw.get("sale_price"))
    fba_fee = to_float(raw.get("fba_fee"))
    purchase_cost_cny = to_float(raw.get("purchase_cost_cny"))
    inbound_placement_fee_cny = to_float(raw.get("inbound_placement_fee_cny"))
    first_leg_shipping_cny = resolve_first_leg_shipping_cny(raw)
    return_rate = to_float(raw.get("return_rate"))
    if return_rate is None:
        return_rate = default_return_rate(package)
    return {
        "sale_price": sale_price,
        "purchase_cost_cny": purchase_cost_cny,
        "purchase_cost": cny_to_site_currency(purchase_cost_cny, exchange_rate),
        "exchange_rate": exchange_rate,
        "fba_fee": fba_fee,
        "first_leg_shipping_cny": first_leg_shipping_cny,
        "first_leg_shipping": cny_to_site_currency(first_leg_shipping_cny, exchange_rate),
        "actual_weight_g": to_float(raw.get("actual_weight_g")),
        "volume_weight_g": to_float(raw.get("volume_weight_g")),
        "billable_weight_g": billable_weight(raw),
        "first_leg_channel": compact_text(raw.get("first_leg_channel")),
        "inbound_placement_fee_cny": inbound_placement_fee_cny,
        "inbound_placement_fee": cny_to_site_currency(inbound_placement_fee_cny, exchange_rate),
        "commission_rate": to_float(raw.get("commission_rate")) if to_float(raw.get("commission_rate")) is not None else DEFAULT_COMMISSION_RATE,
        "storage_fee_rate": to_float(raw.get("storage_fee_rate")) if to_float(raw.get("storage_fee_rate")) is not None else DEFAULT_STORAGE_FEE_RATE,
        "ad_rate_assumption": to_float(raw.get("ad_rate")) if to_float(raw.get("ad_rate")) is not None else DEFAULT_AD_RATE,
        "return_rate_assumption": return_rate,
        "operator_notes": compact_text(raw.get("operator_notes")),
    }


def validate_profit_inputs(inputs: dict[str, Any]) -> list[str]:
    required = [
        ("sale_price", "建议售价"),
        ("purchase_cost_cny", "采购价"),
        ("exchange_rate", "站点汇率"),
        ("fba_fee", "FBA费用"),
        ("first_leg_shipping_cny", "头程费用"),
        ("inbound_placement_fee_cny", "入库配置费"),
    ]
    missing = [label for key, label in required if inputs.get(key) is None]
    if inputs.get("exchange_rate") == 0:
        missing.append("站点汇率不能为 0")
    return missing


def build_profit_reference(inputs: dict[str, Any], missing: list[str]) -> dict[str, Any]:
    if missing:
        return {
            "status": "待补",
            "missing_fields": missing,
            "base_fba_gross_profit": "待补",
            "base_fba_margin": "待补",
            "post_ads_returns_gross_profit": "待补",
            "post_ads_returns_margin": "待补",
            "input_snapshot": inputs,
        }
    profit_inputs = ProfitInputs(
        sale_price=float(inputs["sale_price"]),
        purchase_cost=float(inputs["purchase_cost"]),
        fba_fee=float(inputs["fba_fee"]),
        inbound_placement_fee=float(inputs["inbound_placement_fee"]),
        first_leg_shipping=float(inputs["first_leg_shipping"]),
        commission_rate=float(inputs["commission_rate"]),
        storage_fee_rate=float(inputs["storage_fee_rate"]),
        ad_rate=float(inputs["ad_rate_assumption"]),
        return_rate=float(inputs["return_rate_assumption"] or 0.0),
    )
    breakdown = calc_profit_breakdown(profit_inputs)
    return {
        "status": "已计算",
        "base_fba_gross_profit": round_money(breakdown["base_fba_gross_profit"]),
        "base_fba_margin": round_rate(breakdown["base_fba_margin"]),
        "post_ads_returns_gross_profit": round_money(breakdown["post_ads_returns_gross_profit"]),
        "post_ads_returns_margin": round_rate(breakdown["post_ads_returns_margin"]),
        "cost_breakdown": {key: round_money(value) for key, value in breakdown.items() if key not in {"base_fba_margin", "post_ads_returns_margin"}},
        "rate_assumptions": {
            "commission_rate": inputs["commission_rate"],
            "storage_fee_rate": inputs["storage_fee_rate"],
            "ad_rate": inputs["ad_rate_assumption"],
            "return_rate": inputs["return_rate_assumption"],
        },
        "input_snapshot": inputs,
        "notes": "利润只做参考；FBA、采购价、头程和入库配置费以运营复核为准。",
    }


def update_decision_review(decision: dict[str, Any], missing: list[str], profit_reference: dict[str, Any]) -> dict[str, Any]:
    updated = dict(decision or {})
    facts = list(updated.get("facts", []))
    inferences = list(updated.get("inferences", []))
    action_items = list(updated.get("action_items", []))
    currency_code = profit_reference.get("currency_code", "USD")
    missing_inputs = [item for item in updated.get("missing_inputs", []) if item not in {"建议售价", "采购价", "FBA费用", "头程费用", "入库配置费"}]
    if missing:
        missing_inputs.extend(item for item in missing if item not in missing_inputs)
        action_items = replace_profit_action(action_items, "补齐利润复核字段：" + "、".join(missing) + "。")
    else:
        facts.append(
            "利润复核：基础 FBA 毛利 "
            f"{format_money(profit_reference.get('base_fba_gross_profit'), currency_code)}，毛利率 {format_rate(profit_reference.get('base_fba_margin'))}；"
            "扣广告和退货后 FBA 毛利 "
            f"{format_money(profit_reference.get('post_ads_returns_gross_profit'), currency_code)}，毛利率 {format_rate(profit_reference.get('post_ads_returns_margin'))}。"
        )
        action_items = replace_profit_action(action_items, "结合利润复核结果，继续复核知产/合规和供应链打样可行性。")
        post_margin = profit_reference.get("post_ads_returns_margin")
        if isinstance(post_margin, (int, float)) and post_margin < 0:
            inferences.append("扣广告和退货后毛利率为负，当前报价结构不适合继续推进，除非采购价或费用有明显优化空间。")
    updated["facts"] = facts
    updated["inferences"] = inferences
    updated["missing_inputs"] = missing_inputs
    updated["action_items"] = action_items
    updated["risk_matrix"] = update_profit_risk(updated.get("risk_matrix", []), missing, profit_reference)
    return updated


def replace_profit_action(action_items: list[str], replacement: str) -> list[str]:
    filtered = [item for item in action_items if "利润" not in item and "FBA费用" not in item and "采购价" not in item]
    return [replacement] + filtered


def update_profit_risk(risks: list[dict[str, Any]], missing: list[str], profit_reference: dict[str, Any]) -> list[dict[str, Any]]:
    updated: list[dict[str, Any]] = []
    replaced = False
    for item in risks:
        if item.get("dimension") == "利润不确定性":
            updated.append(profit_risk_item(missing, profit_reference))
            replaced = True
        else:
            updated.append(item)
    if not replaced:
        updated.append(profit_risk_item(missing, profit_reference))
    return updated


def profit_risk_item(missing: list[str], profit_reference: dict[str, Any]) -> dict[str, str]:
    currency_code = profit_reference.get("currency_code", "USD")
    if missing:
        return {
            "dimension": "利润不确定性",
            "level": "待补",
            "basis": "仍缺少：" + "、".join(missing),
            "next_check": "补齐利润模板后再计算毛利率。",
        }
    post_margin = profit_reference.get("post_ads_returns_margin")
    level = "高" if isinstance(post_margin, (int, float)) and post_margin < 0.05 else "中"
    return {
        "dimension": "利润不确定性",
        "level": level,
        "basis": (
            f"基础 FBA 毛利 {format_money(profit_reference.get('base_fba_gross_profit'), currency_code)}，"
            f"扣广告和退货后毛利 {format_money(profit_reference.get('post_ads_returns_gross_profit'), currency_code)}，"
            f"扣广告和退货后毛利率 {format_rate(profit_reference.get('post_ads_returns_margin'))}"
        ),
        "next_check": "用亚马逊后台/收入计算器和供应链报价做最终复核。",
    }


def update_status_card(status: dict[str, Any], missing: list[str]) -> dict[str, Any]:
    updated = dict(status or {})
    if missing:
        updated["next_step"] = "补齐利润复核字段：" + "、".join(missing)
    else:
        updated["next_step"] = "结合利润复核结果、知产/合规和供应链能力判断是否进入打样。"
    return updated


def resolve_first_leg_shipping_cny(raw: dict[str, Any]) -> float | None:
    direct = to_float(raw.get("first_leg_shipping_cny"))
    if direct is not None:
        return direct
    actual_weight_g = to_float(raw.get("actual_weight_g"))
    volume_weight_g = to_float(raw.get("volume_weight_g"))
    channel = compact_text(raw.get("first_leg_channel"))
    if actual_weight_g is None or volume_weight_g is None or not channel:
        return None
    rate = FIRST_LEG_RATE_CNY_PER_KG.get(channel)
    if rate is None:
        return None
    return calc_first_leg_shipping(calc_billable_weight_g(actual_weight_g, volume_weight_g), rate)


def billable_weight(raw: dict[str, Any]) -> float | None:
    actual_weight_g = to_float(raw.get("actual_weight_g"))
    volume_weight_g = to_float(raw.get("volume_weight_g"))
    if actual_weight_g is None or volume_weight_g is None:
        return None
    return calc_billable_weight_g(actual_weight_g, volume_weight_g)


def cny_to_site_currency(value: float | None, exchange_rate: float | None) -> float | None:
    if value is None or exchange_rate in (None, 0):
        return None
    return value / exchange_rate


def default_return_rate(package: dict[str, Any]) -> float:
    return_risk = package.get("return_risk", {})
    for key in ("category_return_rate", "market_return_rate"):
        value = to_float(return_risk.get(key))
        if value is not None:
            return value
    return 0.0


def to_float(value: Any) -> float | None:
    if value in (None, "", "待补"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if "%" in str(value):
        return number / 100
    return number


def clean_value(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def compact_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def round_money(value: float) -> float:
    return round(value, 2)


def round_rate(value: float) -> float:
    return round(value, 4)


def format_rate(value: Any) -> str:
    if isinstance(value, (int, float)):
        return f"{value * 100:.2f}%"
    return str(value)


def site_currency_code(site: Any) -> str:
    text = compact_text(site).upper()
    mapping = {
        "US": "USD",
        "CA": "CAD",
        "UK": "GBP",
        "EU": "EUR",
        "DE": "EUR",
        "FR": "EUR",
        "IT": "EUR",
        "ES": "EUR",
        "JP": "JPY",
        "AU": "AUD",
        "MX": "MXN",
        "USD": "USD",
        "CAD": "CAD",
        "GBP": "GBP",
        "EUR": "EUR",
        "JPY": "JPY",
        "AUD": "AUD",
        "MXN": "MXN",
    }
    return mapping.get(text, text or "USD")


def format_money(value: Any, currency_code: str) -> str:
    if isinstance(value, (int, float)):
        formatted = f"{abs(value):,.2f}" if abs(value) >= 1000 else f"{abs(value):.2f}"
        sign = "-" if value < 0 else ""
        return f"{currency_code} {sign}{formatted}"
    return str(value)


if __name__ == "__main__":
    raise SystemExit(main())
